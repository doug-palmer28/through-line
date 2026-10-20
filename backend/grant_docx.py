"""The grant application as an editable Word document (numbers already filled with the latest data)."""

import io
import re
from datetime import date

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, RGBColor

PLACEHOLDER = re.compile(r"(\[ADD:[^\]]*\])")


def _rgb(hex_color: str) -> RGBColor:
    h = (hex_color or "#2B2B2B").lstrip("#")
    return RGBColor.from_string(h.upper() if len(h) == 6 else "2B2B2B")


def _para(doc, text: str, size: int = 11, italic: bool = False, color: RGBColor | None = None):
    """A paragraph where [ADD: …] placeholders are highlighted so they can't be missed."""
    p = doc.add_paragraph()
    for piece in PLACEHOLDER.split(text):
        if not piece:
            continue
        run = p.add_run(piece)
        run.font.size, run.italic = Pt(size), italic
        if color:
            run.font.color.rgb = color
        if PLACEHOLDER.fullmatch(piece):
            run.bold = True
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def build(view: dict, org_name: str, brand_color: str = "") -> bytes:
    d, profile = view["draft"], view.get("profile") or {}
    accent = _rgb(brand_color)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    for level in (0, 1, 2):
        doc.styles["Title" if level == 0 else f"Heading {level}"].font.color.rgb = accent

    doc.add_heading(d["title"], 0)
    _para(doc, f"{org_name} · {view['foundation']}" + (f" · {d['opportunity']}" if d.get("opportunity") else ""), 11,
          color=RGBColor(0x6F, 0x6A, 0x63))
    _para(doc, f"Draft prepared {date.today():%B %d, %Y} with Through Line. Our numbers are filled in from our latest "
               "data. Highlighted [ADD: …] notes mark information to add before submitting.", 9, italic=True,
          color=RGBColor(0x6F, 0x6A, 0x63))

    doc.add_heading("Why we're a strong fit", 1)
    _para(doc, d["fit_summary"])

    for a in d["answers"]:
        doc.add_heading(a["question"], 2)
        limit = f"Word limit: {a['word_limit']} · " if a.get("word_limit") else ""
        flag = " · over the limit, please trim" if a.get("over_limit") else ""
        _para(doc, f"{limit}This draft: {a.get('words', 0)} words{flag}", 9, italic=True,
              color=RGBColor(0xB0, 0x3A, 0x26) if a.get("over_limit") else RGBColor(0x6F, 0x6A, 0x63))
        for block in [b for b in a["answer"].split("\n\n") if b.strip()]:
            _para(doc, block.strip())

    doc.add_heading("Before you submit", 1)
    if d.get("checklist"):
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light List Accent 1"
        for cell, text in zip(table.rows[0].cells, ("Item", "Status", "What to do")):
            cell.text = text
        for item in d["checklist"]:
            row = table.add_row().cells
            row[0].text, row[1].text = item["item"], "Ready" if item["status"] == "ready" else "Needs input"
            row[2].text = item.get("note", "")
    for need in d.get("needs_input", []):
        _para(doc, f"[ADD: {need}]" if not need.startswith("[ADD") else need)

    if profile:
        doc.add_heading(f"What {profile.get('foundation', view['foundation'])} cares about", 1)
        _para(doc, profile.get("mission", ""), italic=True)
        for pr in profile.get("priorities", []):
            ids = f" [{', '.join(pr.get('source_ids', []))}]" if pr.get("source_ids") else ""
            doc.add_paragraph(f"{pr['priority']}: {pr.get('evidence', '')}{ids}", style="List Bullet")
        if profile.get("sources"):
            doc.add_heading("Sources", 2)
            for s in profile["sources"]:
                status = "verified in search" if s.get("verified") else "not verified, please check"
                doc.add_paragraph(f"[{s['id']}] {s['title']}, {s['publisher']}"
                                  + (f". {s['url']}" if s.get("url") else "") + f" ({status})", style="List Number")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
