"""Logic model + measurement plan -> an editable PowerPoint (built in memory).

Clients edit in the app; this deck is the take-away they can keep editing in
PowerPoint and share with their board or funders.
"""

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from models import LogicModel, MeasurementPlan

# Softened Bauhaus palette (matches the app)
RED, BLUE, YELLOW = RGBColor(0xD9, 0x6C, 0x4F), RGBColor(0x3E, 0x63, 0xB0), RGBColor(0xE8, 0xB8, 0x4A)
BLACK, CREAM, WHITE = RGBColor(0x2B, 0x2B, 0x2B), RGBColor(0xF7, 0xF4, 0xEE), RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x4E, 0x8C, 0x7A)
SCALE_LABELS = {
    "agree_5": "Strongly disagree → Strongly agree (1–5)",
    "confidence_5": "Not at all → Extremely confident (1–5)",
    "frequency_5": "Never → Always (1–5)",
    "scale_0_10": "0–10 rating",
    "yes_no": "Yes / No",
    "number": "Number",
    "open_text": "Short written answer",
}
METHOD_LABELS = {"participant_survey": "Participant survey", "staff_assessment": "Staff assessment",
                 "program_records": "Program records"}
FONT = "Century Gothic"

STAGES = [  # (field, header, header fill, header text color)
    ("inputs", "Inputs", BLACK, WHITE),
    ("activities", "Activities", RED, WHITE),
    ("outputs", "Outputs", BLUE, WHITE),
    ("short_term_outcomes", "Short-term outcomes", YELLOW, BLACK),
    ("long_term_outcomes", "Long-term outcomes", RED, WHITE),
]
BULLET = "• "


# ============================================================ building


def _background(slide, color=CREAM):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _text(slide, left, top, width, height, text, *, size=18, bold=False, color=BLACK, name=None, upper=False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line.upper() if upper else line
        for run in p.runs:
            run.font.size, run.font.bold, run.font.name = Pt(size), bold, FONT
            run.font.color.rgb = color
        p.space_after = Pt(6)
    return box


def _shape(slide, kind, left, top, size, color):
    s = slide.shapes.add_shape(kind, Inches(left), Inches(top), Inches(size), Inches(size))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def _bauhaus_mark(slide, left, top, size):
    _shape(slide, MSO_SHAPE.OVAL, left, top, size, RED)
    _shape(slide, MSO_SHAPE.RECTANGLE, left + size * 1.1, top, size, BLUE)
    _shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, left + size * 2.2, top, size, YELLOW)


def _title(slide, text, accent):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.35), Inches(7.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    _text(slide, 0.8, 0.45, 9.5, 0.9, text, size=30, bold=True, upper=True)
    _bauhaus_mark(slide, 11.3, 0.55, 0.42)


def _bullets(items):
    return [BULLET + i for i in items] or [BULLET]


def _cell(cell, text, *, size=11, bold=False, fill=WHITE, color=BLACK):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.text = text
    cell.text_frame.word_wrap = True
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(size), bold, FONT, color


def _plan_slides(prs, blank, plan: MeasurementPlan) -> None:
    # Overview table: one row per metric
    s = prs.slides.add_slide(blank)
    _background(s, WHITE)
    _title(s, "Measurement plan", GREEN)
    _text(s, 0.8, 1.25, 11.8, 0.7, plan.summary, size=13)
    rows = len(plan.metrics) + 1
    shape = s.shapes.add_table(rows, 5, Inches(0.6), Inches(2.05), Inches(12.1), Inches(0.5 * rows))
    table = shape.table
    widths = [2.2, 2.9, 1.9, 2.3, 2.8]
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    for col, head in enumerate(["Metric", "Outcome it tracks", "Method", "When", "How it's scored"]):
        _cell(table.cell(0, col), head.upper(), size=11, bold=True, fill=BLACK, color=WHITE)
    for row, m in enumerate(plan.metrics, start=1):
        for col, text in enumerate([m.name, m.outcome, METHOD_LABELS[m.method], m.timing, m.how_scored]):
            _cell(table.cell(row, col), text, size=10, bold=(col == 0))

    # One slide per tool, listing its questions and scales
    for tool in plan.tools:
        s = prs.slides.add_slide(blank)
        _background(s)
        _title(s, tool.name, BLUE)
        _text(s, 0.8, 1.25, 11.8, 0.5, f"{METHOD_LABELS[tool.method]} · {tool.when}", size=13, color=GREEN)
        lines = []
        for m in plan.metrics:
            if m.name not in tool.metrics:
                continue
            lines.append(m.name.upper())
            lines.extend(
                f"{BULLET}{q.text}   [{SCALE_LABELS[q.scale]}]" + (f"  ({', '.join(q.source_ids)})" if q.source_ids else "")
                for q in m.questions if q.included
            )
        _text(s, 0.8, 1.9, 11.8, 5.2, lines or ["(No metrics assigned to this tool yet.)"], size=13)


def _sources_slides(prs, blank, alignment: list[dict], sources: list[dict]) -> None:
    if alignment:
        s = prs.slides.add_slide(blank)
        _background(s)
        _title(s, "How this lines up with funders", YELLOW)
        _text(s, 0.8, 1.6, 11.8, 5.4, [
            f"{BULLET}{a['note']}" + (f"  ({', '.join(a.get('source_ids', []))})" if a.get("source_ids") else "")
            for a in alignment
        ], size=15)
    if sources:
        s = prs.slides.add_slide(blank)
        _background(s, WHITE)
        _title(s, "Sources", GREEN)
        lines = []
        for c in sources:
            status = "verified link" if c.get("verified") else "unverified: check before citing"
            lines.append(f"{c['id']}  {c['title']} ({c['publisher']}). {status}")
            if c.get("url"):
                lines.append(f"      {c['url']}")
        _text(s, 0.8, 1.5, 11.8, 5.6, lines, size=11)


def build_deck(model: LogicModel, org_name: str, version: int, plan: MeasurementPlan | None = None,
               alignment: list[dict] | None = None, sources: list[dict] | None = None) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Title
    s = prs.slides.add_slide(blank)
    _background(s)
    _shape(s, MSO_SHAPE.OVAL, 8.4, 0.9, 3.6, RED)
    _shape(s, MSO_SHAPE.RECTANGLE, 10.2, 3.6, 2.6, BLUE)
    _shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, 7.6, 4.2, 2.6, YELLOW)
    _text(s, 0.8, 1.6, 7.2, 0.5, "LOGIC MODEL", size=16, bold=True)
    _text(s, 0.8, 2.1, 7.2, 1.6, org_name, size=40, bold=True, upper=True)
    _text(s, 0.8, 3.9, 7.0, 0.8, model.program_name, size=22, name="TL_program_name")
    _text(s, 0.8, 6.5, 7.5, 0.5, f"Version {version} · Made with Through Line",
          size=11, color=RGBColor(0x55, 0x55, 0x55))

    # 2. Theory of change
    s = prs.slides.add_slide(blank)
    _background(s)
    _title(s, "Theory of change", RED)
    _text(s, 0.8, 2.0, 11.6, 4.0, model.theory_of_change, size=26, name="TL_theory_of_change")

    # 3. Logic model table
    s = prs.slides.add_slide(blank)
    _background(s, WHITE)
    _title(s, "Logic model", BLUE)
    table_shape = s.shapes.add_table(2, 5, Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.4))
    table_shape.name = "TL_logic_model"
    table = table_shape.table
    table.rows[0].height = Inches(0.6)
    table.rows[1].height = Inches(4.8)
    for col, (field, header, fill, fg) in enumerate(STAGES):
        head = table.cell(0, col)
        head.fill.solid()
        head.fill.fore_color.rgb = fill
        head.text = header.upper()
        for p in head.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            for r in p.runs:
                r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(13), True, FONT, fg
        body = table.cell(1, col)
        body.fill.solid()
        body.fill.fore_color.rgb = WHITE
        tf = body.text_frame
        tf.word_wrap = True
        for i, item in enumerate(_bullets(getattr(model, field))):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = item
            p.space_after = Pt(6)
            for r in p.runs:
                r.font.size, r.font.name, r.font.color.rgb = Pt(12), FONT, BLACK

    # 4. Assumptions and external factors
    s = prs.slides.add_slide(blank)
    _background(s)
    _title(s, "Assumptions & external factors", YELLOW)
    _text(s, 0.8, 1.7, 5.6, 0.5, "Assumptions", size=16, bold=True, color=RED, upper=True)
    _text(s, 0.8, 2.3, 5.6, 4.5, _bullets(model.assumptions), size=16, name="TL_assumptions")
    _text(s, 7.0, 1.7, 5.6, 0.5, "External factors", size=16, bold=True, color=BLUE, upper=True)
    _text(s, 7.0, 2.3, 5.6, 4.5, _bullets(model.external_factors), size=16, name="TL_external_factors")

    # 5+. Measurement plan
    if plan:
        _plan_slides(prs, blank, plan)

    # Alignment and sources (logic model sources + plan sources, de-duplicated by id)
    all_sources = {c["id"]: c for c in (sources or [])}
    if plan:
        all_sources.update({c.id: c.model_dump() for c in plan.sources})
    _sources_slides(prs, blank, alignment or [], sorted(all_sources.values(), key=lambda c: int(c["id"][1:] or 0)
                                                       if c["id"][1:].isdigit() else 0))

    # Closing
    s = prs.slides.add_slide(blank)
    _background(s, BLACK)
    _text(s, 0.8, 0.6, 11, 0.9, "This deck is yours", size=30, bold=True, color=WHITE, upper=True)
    _text(s, 0.8, 1.8, 11.5, 4.5, [
        "Edit anything: it's a normal PowerPoint file.",
        "Next in Through Line: build these questions into surveys and staff forms, collect pre and post data, "
        "and turn the results into charts for funders.",
    ], size=18, color=WHITE)
    _bauhaus_mark(s, 0.8, 6.2, 0.5)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
