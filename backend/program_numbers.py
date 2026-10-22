"""The monthly program-numbers workbook: download a template, fill it in, upload it back.

The template lists every number the organization tracks (from its measurement plan plus any custom
ones), one row each, with the month already filled in. Staff type the values and upload the file;
each row then replaces that number's value for that month, so re-uploading a corrected file is safe.
Uploads are read in memory and never saved.
"""

import io
import math
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from reports import excel_safe
from security import check_zip

# A sheet can claim a million rows; only read as far as a real template could go.
MAX_SHEET_ROWS = 2000
MAX_SHEET_COLS = 200

HEADER = ["Number id", "What to enter", "Unit", "How it's counted", "Month (YYYY-MM)", "Value", "Note (optional)"]
INK, FILL, INPUT = "23262B", "E9EEF7", "FFF8E6"


def template_xlsx(org_name: str, indicators: list, entries: list, month: str) -> bytes:
    """indicators: Indicator rows; entries: IndicatorEntry rows (used to pre-fill this month's values)."""
    current = {(e.indicator_id, e.period.strftime("%Y-%m")): e for e in entries}
    wb = Workbook()
    ws = wb.active
    ws.title = "Program numbers"
    ws.append([excel_safe(f"{org_name}: program numbers for {month}")])
    ws["A1"].font = Font(bold=True, size=14, color=INK)
    ws.append(["Fill in the Value column (yellow), save, and upload this file on the Program numbers tab. "
               "Leave a value blank to skip it. Don't change the Number id column."])
    ws["A2"].font = Font(italic=True, color="6B6F76")
    ws.append([])
    ws.append(HEADER)
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for ind in indicators:
        e = current.get((ind.id, month))
        ws.append([ind.id, excel_safe(ind.name), excel_safe(ind.unit), excel_safe(ind.description), month,
                   e.value if e else None, excel_safe(e.note) if e else ""])  # staff text can't become a formula
        row = ws.max_row
        ws.cell(row, 6).fill = PatternFill("solid", fgColor=INPUT)
        for col in (2, 4):
            ws.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
    last = ws.max_row
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        error="Please enter a number (0 or more).", errorTitle="Numbers only")
    ws.add_data_validation(dv)
    dv.add(f"F5:F{max(last, 5)}")
    for col, width in zip("ABCDEFG", (11, 60, 12, 50, 16, 14, 30)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "B5"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _month(value) -> str | None:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m")
    text = str(value or "").strip()
    m = re.fullmatch(r"(\d{4})-(\d{1,2})(?:-\d{1,2})?", text)
    if m and 1 <= int(m.group(2)) <= 12:
        return f"{m.group(1)}-{int(m.group(2)):02d}"
    return None


def parse_upload(data: bytes) -> tuple[list[dict], list[dict]]:
    """Return (rows, errors). Each row: {indicator_id, month, value, note, line}."""
    try:
        check_zip(data)
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("That file couldn't be opened as an Excel workbook (.xlsx).") from exc
    ws = wb["Program numbers"] if "Program numbers" in wb.sheetnames else wb.worksheets[0]
    rows, errors, header_seen = [], [], False
    for i, row in enumerate(ws.iter_rows(max_row=MAX_SHEET_ROWS, max_col=MAX_SHEET_COLS, values_only=True), start=1):
        cells = list(row) + [None] * 7
        if not header_seen:
            header_seen = str(cells[0] or "").strip().lower() == "number id"
            continue
        ind_id, name, _unit, _how, month, value, note = cells[:7]
        if ind_id in (None, "") and value in (None, ""):
            continue
        if value in (None, ""):
            continue  # blank = skip
        try:
            ind_id = int(ind_id)
        except (TypeError, ValueError):
            errors.append({"line": i, "error": "The Number id is missing or was changed."})
            continue
        ym = _month(month)
        if not ym:
            errors.append({"line": i, "error": f"'{name}': the month should look like 2026-10."})
            continue
        try:
            num = float(value)
        except (TypeError, ValueError):
            errors.append({"line": i, "error": f"'{name}': the value should be a number."})
            continue
        if not math.isfinite(num) or num < 0 or num > 1_000_000_000:  # "nan"/"inf" would break every page
            errors.append({"line": i, "error": f"'{name}': the value is out of range."})
            continue
        rows.append({"indicator_id": ind_id, "month": ym, "value": num, "note": str(note or "")[:300], "line": i})
    if not header_seen:
        raise ValueError("This doesn't look like the Through Line program-numbers template. Download a fresh one.")
    return rows, errors
