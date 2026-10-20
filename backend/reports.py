"""Dashboard numbers and Excel downloads (one file per chart, plus all data).

The PowerPoint/PDF designs live in design.py; this module only computes the numbers.

Privacy rules applied here:
  * Demographic groups with fewer than MIN_GROUP people are folded into "Groups under 5 (hidden)",
    and a demographic chart with fewer than MIN_GROUP people in total isn't shown at all.
  * Downloads use participant codes, never names (we don't have names).
  * Text written by participants is escaped in Excel so it can't run as a formula.
"""

import io
import re
from collections import Counter, defaultdict
from statistics import mean
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from sqlmodel import Session, select

import surveys
from db import Indicator, IndicatorEntry, Instrument, Organization, Participant, Response

MIN_GROUP = 5
DEMO_TITLES = {
    "age_range": "Age", "gender": "Gender", "race_ethnicity": "Race and ethnicity",
    "primary_language": "Language spoken at home", "town": "Town", "household": "Children under 18 at home",
}
IMPROVE_BY = {"agree_5": 1, "confidence_5": 1, "frequency_5": 1, "scale_0_10": 1}  # points that count as improving
HIDDEN = "Groups under 5 (hidden)"


# ============================================================ computing


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", str(text).lower()).strip("-")[:48] or "item"


def chart_index(data: dict) -> dict[str, dict]:
    """Charts by id, including older ids (aliases) so saved reports keep finding their charts."""
    index = {}
    for c in data["charts"]:
        for key in [*c.get("aliases", []), c["id"]]:
            index[key] = c
    return index


def _improved(scale: str, pre: float, post: float) -> bool:
    return post - pre >= IMPROVE_BY.get(scale, 1e-9)


def compute(db: Session, org: Organization) -> dict:
    instruments = db.exec(select(Instrument).where(Instrument.org_id == org.id)).all()
    responses = db.exec(select(Response).where(Response.org_id == org.id)).all()
    participants = db.exec(select(Participant).where(Participant.org_id == org.id, Participant.active == True)).all()  # noqa: E712
    by_inst: dict[int, list[Response]] = defaultdict(list)
    for r in responses:
        by_inst[r.instrument_id].append(r)

    charts, metric_rows, open_answers = [], [], []
    for inst in instruments:
        rs = by_inst.get(inst.id, [])
        qmap = {q["id"]: q for q in inst.questions}
        # Score only the questions currently shown, so pre and post always compare the same items.
        scored_qs = [q for q in inst.questions if not q["hidden"]]
        scores = {(r.participant_id, r.wave): surveys.metric_score(scored_qs, r.answers) for r in rs}
        metrics = sorted({q["metric"] for q in scored_qs if q.get("metric") and q["scale"] in surveys.SCORED})
        for metric in metrics:
            scale = surveys.metric_scale(scored_qs, metric)
            pairs = [(scores[(pid, "pre")][metric], scores[(pid, "post")][metric])
                     for (pid, wave) in scores if wave == "pre"
                     if (pid, "post") in scores and metric in scores[(pid, "pre")] and metric in scores[(pid, "post")]]
            pre_all = [s[metric] for (pid, w), s in scores.items() if w == "pre" and metric in s]
            post_all = [s[metric] for (pid, w), s in scores.items() if w == "post" and metric in s]
            row = {"metric": metric, "instrument": inst.name, "scale": scale or "", "n_pre": len(pre_all),
                   "n_post": len(post_all), "n_matched": len(pairs),
                   "mean_pre": round(mean(p for p, _ in pairs), 2) if pairs else None,
                   "mean_post": round(mean(q for _, q in pairs), 2) if pairs else None,
                   "pct_improved": round(100 * sum(_improved(scale or "", p, q) for p, q in pairs) / len(pairs))
                   if pairs and scale != "mixed" else None}
            metric_rows.append(row)
            if pairs and scale != "mixed":
                yes_no = scale == "yes_no"
                unit = "%" if yes_no else surveys.SCALES.get(scale, {}).get("label", "")
                if any(q.get("reverse") for q in scored_qs if q.get("metric") == metric):
                    unit += " · reverse-scored, so higher = better"
                vals = [round(row["mean_pre"] * 100), round(row["mean_post"] * 100)] if yes_no else [row["mean_pre"], row["mean_post"]]
                charts.append({
                    # Named by metric (not position) so ids stay the same as more data arrives;
                    # the old position-based id is kept as an alias for reports saved before.
                    "id": f"metric-{inst.id}-{slug(metric)}", "aliases": [f"metric-{inst.id}-{len(charts)}"],
                    "group": "outcomes", "metric": metric,
                    "title": f"{metric}: {'share answering yes' if yes_no else 'average'}, before and after",
                    "kind": "bar", "unit": unit, "categories": ["Before (pre)", "After (post)"],
                    "series": [{"name": "% yes" if yes_no else "Average", "values": vals}],
                    "note": f"{len(pairs)} participants answered both pre and post"
                            + (" (small group: interpret with care)" if len(pairs) < MIN_GROUP else "") + ".",
                })
        for r in rs:
            for qid, val in r.answers.items():
                q = qmap.get(qid)
                if q and q["scale"] == "open_text" and str(val).strip():
                    open_answers.append({"instrument": inst.name, "question": q["text"], "text": str(val),
                                         "wave": r.wave, "at": r.created_at.isoformat()})

    improved = [m for m in metric_rows if m["pct_improved"] is not None]
    if improved:
        charts.insert(0, {
            "id": "improved", "group": "outcomes", "title": "Share of participants who improved", "kind": "bar", "unit": "%",
            "categories": [m["metric"] for m in improved],
            "series": [{"name": "% improved", "values": [m["pct_improved"] for m in improved]}],
            "note": "Among participants with both a pre and a post answer. 'Improved' means at least 1 point higher "
                    "on rating scales, or moving from no to yes.",
        })

    if instruments:
        counts = {(r.instrument_id, r.wave): 0 for r in responses}
        for r in responses:
            counts[(r.instrument_id, r.wave)] += 1
        charts.append({
            "id": "responses", "group": "program", "title": "Surveys completed", "kind": "bar", "unit": "responses",
            "categories": [i.name for i in instruments],
            "series": [{"name": "Pre", "values": [counts.get((i.id, "pre"), 0) for i in instruments]},
                       {"name": "Post", "values": [counts.get((i.id, "post"), 0) for i in instruments]}],
            "note": "Pre and post responses received for each survey tool.",
        })

    # Demographics: one record per participant (their latest answer for each item), small groups hidden.
    latest: dict[int, dict] = {}
    for r in sorted(responses, key=lambda r: r.created_at):
        if r.demographics:
            latest.setdefault(r.participant_id, {}).update(r.demographics)
    keys = [k for k in surveys.DEMOGRAPHICS if any(k in d for d in latest.values())]
    for key in keys:
        tally: Counter = Counter()
        people = 0
        for d in latest.values():
            if key in d:
                people += 1
                for v in (d[key] if isinstance(d[key], list) else [d[key]]):
                    tally[v] += 1
        if people < MIN_GROUP:
            continue
        shown = {k: v for k, v in tally.items() if v >= MIN_GROUP}
        hidden = sum(v for k, v in tally.items() if v < MIN_GROUP)
        cats = [o for o in surveys.DEMOGRAPHICS[key]["options"] if o in shown] + ([HIDDEN] if hidden else [])
        counts = [shown.get(c, hidden if c == HIDDEN else 0) for c in cats]
        multi = surveys.DEMOGRAPHICS[key]["multi"]
        charts.append({
            "id": f"demo-{key}", "group": "people", "title": DEMO_TITLES.get(key, surveys.DEMOGRAPHICS[key]["label"]),
            "kind": "bar", "unit": "%", "categories": cats,
            "series": [{"name": "% of participants", "values": [round(100 * c / people) for c in counts]}],
            "counts": counts,
            "note": f"Share of the {people} participants who answered"
                    + (" (people could choose more than one, so shares can add up to more than 100%)" if multi else "")
                    + f". Groups with fewer than {MIN_GROUP} people are combined to protect privacy.",
        })

    # Program numbers, by month.
    inds = db.exec(select(Indicator).where(Indicator.org_id == org.id)).all()
    entries = db.exec(select(IndicatorEntry).where(IndicatorEntry.org_id == org.id)).all()
    indicator_rows = []
    for ind in inds:
        mine = [e for e in entries if e.indicator_id == ind.id]
        by_month: dict[str, float] = defaultdict(float)
        for e in mine:
            by_month[e.period.strftime("%Y-%m")] += e.value
        total = sum(by_month.values())
        indicator_rows.append({"name": ind.name, "unit": ind.unit, "total": total, "entries": len(mine)})
        if by_month:
            months = sorted(by_month)
            charts.append({
                "id": f"indicator-{ind.id}", "group": "program", "title": ind.name, "kind": "line" if len(months) > 2 else "bar",
                "unit": ind.unit or "count", "categories": months,
                "series": [{"name": ind.name, "values": [by_month[m] for m in months]}],
                # Dollars add up across months; counts like "active landlords" are a level, so show the latest.
                "note": "Entered by staff. " + (f"Total across {len(months)} months: ${total:,.0f}." if ind.unit == "$" else
                        f"Latest month ({months[-1]}): {by_month[months[-1]]:,.0f}{(' ' + ind.unit) if ind.unit else ''}."),
            })

    matched = {r.participant_id for r in responses if r.wave == "pre"} & {r.participant_id for r in responses if r.wave == "post"}
    best = max(improved, key=lambda m: m["n_matched"], default=None)
    return {
        "org_name": org.name,
        "headline": {
            "participants": len(participants),
            "responses_pre": sum(1 for r in responses if r.wave == "pre"),
            "responses_post": sum(1 for r in responses if r.wave == "post"),
            "matched": len(matched),
            "top_improvement": {"metric": best["metric"], "pct": best["pct_improved"]} if best else None,
        },
        "metrics": metric_rows,
        "indicators": indicator_rows,
        "charts": charts,
        "open_responses": sorted(open_answers, key=lambda a: a["at"], reverse=True)[:12],
        "min_group_size": MIN_GROUP,
    }


# ============================================================ Excel


def excel_safe(value: Any) -> Any:
    """Stop text that starts with = + - @ from being treated as a spreadsheet formula."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def chart_xlsx(chart: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Chart data"
    ws.append([excel_safe(chart["title"])])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([excel_safe(chart.get("note", ""))])
    ws.append([])
    ws.append(["Category"] + [excel_safe(s["name"]) for s in chart["series"]])
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2B2B2B")
    for i, cat in enumerate(chart["categories"]):
        ws.append([excel_safe(cat)] + [s["values"][i] for s in chart["series"]])
    last = 4 + len(chart["categories"])
    ch = LineChart() if chart["kind"] == "line" else BarChart()
    ch.title = chart["title"]
    ch.y_axis.title = chart.get("unit", "")
    ch.height, ch.width = 9, 18
    data = Reference(ws, min_col=2, max_col=1 + len(chart["series"]), min_row=4, max_row=last)
    cats = Reference(ws, min_col=1, min_row=5, max_row=last)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ws.add_chart(ch, f"{chr(ord('B') + len(chart['series']) + 1)}4")
    ws.column_dimensions["A"].width = 34
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def data_xlsx(db: Session, org: Organization) -> bytes:
    wb = Workbook()
    code_ws = wb.active
    code_ws.title = "Codebook"
    code_ws.append(["Tool", "Question id", "Question", "Scale", "Values", "Metric", "Asked at", "Reverse-scored", "Hidden"])
    instruments = db.exec(select(Instrument).where(Instrument.org_id == org.id)).all()
    codes = {p.id: p.code for p in db.exec(select(Participant).where(Participant.org_id == org.id)).all()}
    for n, inst in enumerate(instruments, start=1):
        for q in inst.questions:
            opts = surveys.SCALES[q["scale"]]["options"]
            values = ", ".join(f"{i + surveys.MIN_VALUE.get(q['scale'], 0)}={o}" for i, o in enumerate(opts)) if opts else q["scale"]
            code_ws.append([excel_safe(inst.name), q["id"], excel_safe(q["text"]), q["scale"], values, excel_safe(q.get("metric") or ""),
                            "/".join(q["waves"]), "yes" if q.get("reverse") else "", "yes" if q["hidden"] else ""])
        ws = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "-", f"{n}. {inst.name}")[:31])  # Excel forbids \ / * ? : [ ]
        demo_keys = inst.demographics
        header = ["Participant code", "Wave", "Entered by", "Date"] + [f"demo: {k}" for k in demo_keys] \
                 + [f"{q['text'][:60]} [{q['id']}]" for q in inst.questions]
        ws.append([excel_safe(h) for h in header])
        for c in ws[1]:
            c.font = Font(bold=True)
        rs = db.exec(select(Response).where(Response.instrument_id == inst.id, Response.org_id == org.id)
                     .order_by(Response.participant_id, Response.wave.desc())).all()
        for r in rs:
            demo = [", ".join(v) if isinstance(v := r.demographics.get(k, ""), list) else v for k in demo_keys]
            ws.append([codes.get(r.participant_id, ""), r.wave, r.entered_by, r.created_at.date().isoformat()]
                      + [excel_safe(v) for v in demo] + [excel_safe(r.answers.get(q["id"], "")) for q in inst.questions])
    ind_ws = wb.create_sheet("Program numbers")
    ind_ws.append(["Number", "Unit", "Period", "Value", "Note"])
    names = {i.id: i for i in db.exec(select(Indicator).where(Indicator.org_id == org.id)).all()}
    for e in db.exec(select(IndicatorEntry).where(IndicatorEntry.org_id == org.id).order_by(IndicatorEntry.period)).all():
        ind = names.get(e.indicator_id)
        ind_ws.append([excel_safe(ind.name if ind else ""), excel_safe(ind.unit if ind else ""), e.period.isoformat(), e.value, excel_safe(e.note)])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
