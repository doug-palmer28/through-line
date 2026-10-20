"""Build a walkthrough deck from a test-profile run: screenshots + real outputs + native charts.

    .venv/Scripts/python scripts/build_demo_deck.py data/test_profiles/keystone_family_housing 2026-10-20_live_demo
"""

import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

profile_dir = Path(sys.argv[1])
run_dir = profile_dir / "runs" / sys.argv[2]
screens = profile_dir / "demo_screens"
profile = json.loads((profile_dir / "profile.json").read_text(encoding="utf-8"))
studio = json.loads((run_dir / "studio.json").read_text(encoding="utf-8"))
tools_log = json.loads((run_dir / "run_log_tools_collect.json").read_text(encoding="utf-8"))
dash = json.loads((run_dir / "dashboard.json").read_text(encoding="utf-8"))
ORG = profile["org"]["name"]

INK, MUTED, PAPER, WHITE, LINE = (RGBColor(0x2B, 0x2B, 0x2B), RGBColor(0x6F, 0x6A, 0x63), RGBColor(0xF7, 0xF4, 0xEE),
                                  RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xE5, 0xDF, 0xD5))
RED, BLUE, YELLOW, GREEN, PURPLE = (RGBColor(0xD9, 0x6C, 0x4F), RGBColor(0x3E, 0x63, 0xB0), RGBColor(0xE8, 0xB8, 0x4A),
                                    RGBColor(0x4E, 0x8C, 0x7A), RGBColor(0x8A, 0x6F, 0xB0))
HEAD, BODY = "Century Gothic", "Segoe UI"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


# ============================================================ helpers


def text(slide, x, y, w, h, lines, size=14, bold=False, color=INK, font=BODY, align=None, space=6):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        runs = line if isinstance(line, list) else [(line, bold, color)]
        for t, b, c in runs:
            r = p.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(size), b, font, c
        p.space_after = Pt(space)
        if align:
            p.alignment = align
    return box


def bullets(slide, x, y, w, h, items, size=14):
    return text(slide, x, y, w, h, [[("•  ", True, MUTED), (it, False, INK)] if isinstance(it, str) else it for it in items],
                size=size, space=8)


def shape(slide, kind, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def marks(slide, x, y, size=0.18):
    shape(slide, MSO_SHAPE.OVAL, x, y, size, size, RED)
    shape(slide, MSO_SHAPE.RECTANGLE, x + size * 1.25, y, size, size, BLUE)
    shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, x + size * 2.5, y, size, size, YELLOW)


def new_slide(title, eyebrow="", step=None, accent=BLUE):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PAPER
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.08, accent)
    if eyebrow:
        text(s, 0.6, 0.3, 9, 0.35, eyebrow.upper(), 10.5, True, MUTED, BODY)
    text(s, 0.6, 0.58, 10.6, 0.8, title, 26, True, INK, HEAD)
    if step:
        chip = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 10.95, 0.32, 1.8, 0.38, WHITE, LINE)
        chip.text_frame.text = step
        for p in chip.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(10), True, BODY, INK
    text(s, 0.6, 7.08, 12, 0.3, f"Live test · {ORG} (fictional) · gpt-6-luna via Portkey · 2026-10-20", 9, False, MUTED)
    marks(s, 12.3, 7.12, 0.14)
    return s


def picture(slide, name, x, y, w=None, h=None, crop_top=0.0, crop_bottom=0.0):
    path = screens / name
    iw, ih = Image.open(path).size
    vis_h = ih * (1 - crop_top - crop_bottom)
    if w and not h:
        h = w * vis_h / iw
    elif h and not w:
        w = h * iw / vis_h
    frame = shape(slide, MSO_SHAPE.RECTANGLE, x - 0.04, y - 0.04, w + 0.08, h + 0.08, WHITE, LINE)
    frame.shadow.inherit = False
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    pic.crop_top, pic.crop_bottom = crop_top, crop_bottom
    return pic


def callout(slide, x, y, w, h, title, body, color=GREEN, size=12.5):
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.07, h, color)
    box = shape(slide, MSO_SHAPE.RECTANGLE, x + 0.07, y, w - 0.07, h, WHITE)
    box.text_frame.word_wrap = True
    text(slide, x + 0.25, y + 0.1, w - 0.4, h - 0.2,
         [[(title, True, INK)], [(body, False, INK)]] if title else [[(body, False, INK)]], size=size, space=4)


def table(slide, x, y, w, rows, col_w, size=11, header_fill=INK):
    t = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.35 * len(rows))).table
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (WHITE if r % 2 else PAPER)
            cell.text = str(val)
            cell.margin_left = cell.margin_right = Inches(0.08)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size, run.font.name = Pt(size), BODY
                    run.font.bold = r == 0
                    run.font.color.rgb = WHITE if r == 0 else INK
    return t


def bar_chart(slide, x, y, w, h, title, cats, series, colors, unit="", horizontal=True, labels=True, max_v=None):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    ch = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    for p in ch.chart_title.text_frame.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(12), True, HEAD, INK
    ch.font.size, ch.font.name = Pt(10), BODY
    ch.has_legend = len(series) > 1
    if ch.has_legend:
        ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = labels
    if labels:
        plot.data_labels.font.size = Pt(10)
        plot.data_labels.number_format = '0"%"' if unit == "%" else "General"
        plot.data_labels.number_format_is_linked = False
    for i, s in enumerate(plot.series):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = colors[i % len(colors)]
    if len(series) == 1 and list(cats) == ["Before (pre)", "After (post)"]:  # color follows the entity: after = terracotta
        pt = plot.series[0].points[1]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = RED
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE
    va.format.line.fill.background()
    if max_v:
        va.maximum_scale = max_v
    va.minimum_scale = 0
    if horizontal:
        ch.category_axis.reverse_order = True
    return ch


def line_chart(slide, x, y, w, h, title, cats, series, colors):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    ch = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    for p in ch.chart_title.text_frame.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(12), True, HEAD, INK
    ch.font.size, ch.font.name = Pt(10), BODY
    ch.has_legend = len(series) > 1
    if ch.has_legend:
        ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    for i, s in enumerate(ch.plots[0].series):
        s.format.line.color.rgb = colors[i % len(colors)]
        s.format.line.width = Pt(2)
        s.smooth = False
    ch.value_axis.major_gridlines.format.line.color.rgb = LINE
    ch.value_axis.minimum_scale = 0
    return ch


def chart_by_title(prefix):
    return next(c for c in dash["charts"] if c["title"].startswith(prefix))


# ============================================================ data pulled from the run

lm = studio["approved_logic_model"]
plan = studio["approved_plan"]
interview = studio["interview"]
review = tools_log["stages"]["tools"]["reviews"][0]
collect = tools_log["stages"]["collect"]
h = dash["headline"]
team_msgs = [t["text"] for t in studio["transcript"] if t["role"] == "team"]
first_draft_msg = next(m for m in team_msgs if m.startswith("Here is the revised logic model"))
lm_feedback = profile["interview"]["logic_model_feedback"]
lm_revision_msg = next(m for m in team_msgs if m.startswith("I added Fresh Start"))
metrics_feedback = profile["interview"]["metrics_feedback"]
metrics_revision_msg = next(m for m in team_msgs if m.startswith("I’ve kept the participant survey"))
sources_lm = studio["logic_model_sources"] or []
TOKENS = {"runs": 28, "requests": 56, "input": 298_248, "output": 40_788, "refused": 3, "audit": 170}

# ============================================================ slides

# 1 Title
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = PAPER
shape(s, MSO_SHAPE.OVAL, 8.6, 1.0, 3.0, 3.0, RED)
shape(s, MSO_SHAPE.RECTANGLE, 10.4, 3.4, 2.3, 2.3, BLUE)
shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, 7.9, 3.9, 2.4, 2.1, YELLOW)
text(s, 0.8, 1.4, 7.2, 0.4, "THROUGH LINE · LIVE WALKTHROUGH", 12, True, MUTED)
text(s, 0.8, 1.9, 7.4, 2.2, "From mission to evidence to story, end to end", 38, True, INK, HEAD)
text(s, 0.8, 4.3, 7.0, 1.4, [f"A full test of the working prototype with a fictional nonprofit, {ORG}, "
                             "using live AI agents (gpt-6-luna via Portkey) and synthetic participant data."], 16, False, MUTED)
text(s, 0.8, 6.4, 8, 0.4, "Prepared for Doug · AI Foundations final project · 2026-10-20", 11, False, MUTED)

# 2 The journey
s = new_slide("What you'll see: one organization, start to finish", "The journey", accent=RED)
steps = [("1", "Sign up", "Staff account first; everything saves to the org", INK),
         ("2", "Share", "Upload what you have (read, never saved)", RED),
         ("3", "Five questions", "AI team interviews the director", YELLOW),
         ("4", "Logic model", "Drafted, peer-reviewed, refined", BLUE),
         ("5", "Metrics", "Research-backed questions & scales", GREEN),
         ("6", "Collect", "Survey tools, portal, staff entry", PURPLE),
         ("7", "Dashboard", "Charts, downloads, report", RED)]
for i, (n, title, sub, col) in enumerate(steps):
    x = 0.6 + i * 1.75
    shape(s, MSO_SHAPE.OVAL, x + 0.5, 1.9, 0.55, 0.55, col)
    text(s, x + 0.5, 1.98, 0.55, 0.4, n, 16, True, WHITE if col != YELLOW else INK, HEAD, PP_ALIGN.CENTER)
    text(s, x, 2.65, 1.6, 0.4, title, 14, True, INK, HEAD, PP_ALIGN.CENTER)
    text(s, x, 3.05, 1.6, 1.0, sub, 11, False, MUTED, BODY, PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        shape(s, MSO_SHAPE.RECTANGLE, x + 1.15, 2.16, 1.15, 0.03, LINE)
callout(s, 0.6, 4.6, 6.0, 2.0, "How this test was run",
        "The Logic Model Studio was driven through the real browser UI so you can see the live team at work. "
        "Survey tools, the Survey Review Bot, 24 participants and program numbers were then loaded by "
        "scripts/run_test_profile.py from the saved test profile.", BLUE)
callout(s, 6.9, 4.6, 5.8, 2.0, "What was real vs. synthetic",
        "Real: every agent decision, question, draft, citation and review (live gpt-6-luna, live web search). "
        "Synthetic: the organization, the director's answers (from the profile's answer bank), and all participant data.", GREEN)

# 3 The test profile
s = new_slide("The saved test profile: reusable fake data", "Test profile", accent=PURPLE)
pp = profile["participants"]
bullets(s, 0.6, 1.5, 6.3, 5.3, [
    f"Organization: {ORG}, family rapid rehousing (fictional, West Haven, CT)",
    "Uploaded document: its 2-page PDF profile from data/test_organizations/",
    f"Staff account: {profile['account']['username']} · {profile['account']['display_name']}",
    f"Interview answer bank: {len(profile['interview']['answer_bank'])} answers, matched to whatever the AI asks",
    "Scripted feedback for one logic model round and one metrics round",
    f"{pp['count']} synthetic participants with fixed demographics and a seeded response pattern "
    f"({int(pp['response_model']['improver_share'] * 100)}% improve; ~8% skip the post survey)",
    "6 months of program numbers, matched to the plan's indicators by keyword, plus 3 custom numbers",
], 14)
callout(s, 7.3, 1.5, 5.4, 2.2, "Where it lives",
        "data/test_profiles/keystone_family_housing/\n  profile.json: the fake firm and its data\n  demo_screens/: screenshots for this deck\n"
        "  runs/<date>/: outputs of each run", PURPLE, 12)
callout(s, 7.3, 4.0, 5.4, 2.6, "Run it again",
        ".venv/Scripts/python scripts/run_test_profile.py data/test_profiles/keystone_family_housing\n\n"
        "Add --from tools (or collect, numbers, export) to resume part-way. Same seed gives the same participant data.",
        GREEN, 12)

# 4 Accounts
s = new_slide("Accounts come first, and they're built to be safe", "Step 1 · Sign up", "Step 1", INK)
picture(s, "02_signup.jpg", 0.6, 1.5, w=7.0)
bullets(s, 8.0, 1.5, 4.8, 5.3, [
    "Signing up creates the organization; the logic model, metrics, surveys and data are all saved to it",
    "Passwords and PINs: Argon2id hashes only",
    "Usernames are never stored, only a keyed hash used to find the account",
    "Sign-in token in an HttpOnly, SameSite=Strict cookie; staff and participants use separate cookies",
    "5 wrong attempts lock the account for 15 minutes",
    "Every query is scoped to the signed-in organization",
], 13)

# 5 Tutorial
s = new_slide("A short tutorial, then a team of five agents", "Step 1 · How it works", "Step 1", RED)
picture(s, "03_tutorial.jpg", 0.6, 1.5, w=7.0)
bullets(s, 8.0, 1.5, 4.8, 5.3, [
    [("Team Lead ● ", True, RED), ("talks with the client and writes every draft", False, INK)],
    [("Sector Expert ▲ ", True, YELLOW), ("practice standards in the client's field", False, INK)],
    [("Client Advocate ■ ", True, BLUE), ("the org's voice, capacity and participants", False, INK)],
    [("Evaluation Specialist ◆ ", True, GREEN), ("logic, measurement, scales", False, INK)],
    [("Funder Researcher ⬢ ", True, PURPLE), ("live web research on what funders want", False, INK)],
    "A stepper across the top shows exactly where you are",
], 13)

# 6 Share
s = new_slide("Share what you already have", "Step 2 · Share", "Step 2", RED)
picture(s, "04_share.jpg", 0.6, 1.5, w=7.0)
bullets(s, 8.0, 1.5, 4.8, 3.2, [
    "Keystone's 2-page PDF profile was uploaded here",
    "Files and website text are read in memory and dropped after intake: never written to disk, never in the audit log",
    "Org name defaults to the signed-in organization",
], 13)
callout(s, 8.0, 4.8, 4.8, 1.8, "Intake took about 15 seconds",
        "The Team Lead read the PDF while four specialists worked in parallel: a practice brief, live funder research "
        "(16 sources, all verified), and proposed interview questions.", GREEN, 12)

# 7 Q1 + huddle
s = new_slide("The interview starts with the team's best question", "Step 3 · Five questions", "Step 3", YELLOW)
picture(s, "05_interview_q1.jpg", 0.6, 1.5, w=7.0)
bullets(s, 8.0, 1.5, 4.8, 5.3, [
    "The welcome shows the team actually read the PDF (HMIS, home visits, furniture)",
    "Question 1 came from the Evaluation Specialist's proposal: what does 'kept their home' mean, and how would you confirm it?",
    "The live huddle (right) shows each agent as a shape and logs every hand-off",
    "\"Why we ask\" explains each question in plain language",
], 13)

# 8 Answer + thinking
s = new_slide("Answering: a word countdown, then the team thinks", "Step 3 · Five questions", "Step 3", YELLOW)
picture(s, "06_interview_answer.jpg", 0.6, 1.55, w=6.0)
picture(s, "07_interview_thinking.jpg", 6.9, 1.55, w=6.0)
text(s, 0.6, 5.65, 6, 0.4, "Answers are capped at 200 words; the bar counts down as you type.", 11, False, MUTED)
text(s, 6.9, 5.65, 6, 0.4, "While the Team Lead works, the huddle switches to \"Working\" and narrates.", 11, False, MUTED)

# 9 Interview table
s = new_slide("All five questions were written live by the AI", "Step 3 · Five questions", "Step 3", YELLOW)
rows = [["#", "What the team asked", "Answer given (profile topic)"]]
for i, (qa, topic) in enumerate(zip(interview, ["how we know", "changes we see", "who we serve", "what we do", "funders"]), 1):
    rows.append([str(i), qa["question"], topic])
table(s, 0.6, 1.5, 12.1, rows, [0.4, 9.2, 2.5], 11)
callout(s, 0.6, 5.55, 12.1, 1.15, "The team adapted",
        "Question 3 asked how case management works, but the test's answer matcher sent the 'who we serve' answer. "
        "The Team Lead noticed the gap and asked about case management again in question 4. (The matcher has since been fixed.)",
        YELLOW, 12)

# 10 Peer review live
s = new_slide("Every draft is peer-reviewed before you see it", "Step 4 · Logic model", "Step 4", BLUE)
picture(s, "08_peer_review_live.jpg", 0.6, 1.5, w=7.0)
bullets(s, 8.0, 1.5, 4.8, 5.3, [
    "The Team Lead writes a working draft",
    "All four specialists review it at the same time: every link lights up",
    "The Team Lead folds their reviews into the version you see",
    "What the reviewers changed (from the live run):",
    "Sector Expert: made housing skills observable; removed 'steadier income'",
    "Client Advocate: case management described as voluntary and tailored",
    "Evaluation Specialist: separated move-in timing from family outcomes",
    "Funder Researcher: centred retention and returns to homelessness",
], 12)

# 11 Draft
s = new_slide("The first draft, ready to edit", "Step 4 · Logic model", "Step 4", BLUE)
picture(s, "09_logic_model_review.jpg", 0.6, 1.5, w=7.0)
callout(s, 8.0, 1.5, 4.8, 3.0, "Theory of change (as drafted)", lm["theory_of_change"], BLUE, 11.5)
bullets(s, 8.0, 4.75, 4.8, 2.0, [
    "Edit any line directly, or ask the team for changes (2 rounds)",
    "\"Approve\" is disabled while edits are unsaved",
], 12)

# 12 Grid
s = new_slide("The logic model workspace", "Step 4 · Logic model", "Step 4", BLUE)
picture(s, "10_logic_model_grid.jpg", 0.6, 1.5, w=7.0)
bullets(s, 8.0, 1.5, 4.8, 5.3, [
    f"Inputs: {len(lm['inputs'])} · Activities: {len(lm['activities'])} · Outputs: {len(lm['outputs'])}",
    f"Short-term outcomes: {len(lm['short_term_outcomes'])} · Long-term: {len(lm['long_term_outcomes'])}",
    "Long-term outcomes use HUD's framing: staying housed after assistance ends; fewer returns to shelter",
    "Each item is an editable line with add/remove",
], 13)

# 13 Alignment
s = new_slide("Research-backed: how it lines up with funders", "Step 4 · Logic model", "Step 4", PURPLE)
picture(s, "11_alignment_funders.jpg", 0.6, 1.5, w=7.0)
bullets(s, 8.0, 1.5, 4.8, 5.3, [
    "The Funder Researcher searched the web live through Portkey; it works",
    "Found real local and federal funders: HUD Continuum of Care, CT Dept. of Housing, City of New Haven ESG, "
    "Community Foundation for Greater New Haven, United Way",
    "Each alignment note carries citation chips (S2, S4, …)",
    "The prompt forbids adding outcomes the client doesn't actually pursue",
], 12.5)

# 14 Sources
s = new_slide("Every source is checked in code", "Step 4 · Logic model", "Step 4", PURPLE)
picture(s, "12_sources_verified.jpg", 0.6, 1.5, w=7.0)
bullets(s, 8.0, 1.5, 4.8, 5.3, [
    "16 sources at intake; all 16 'verified link'",
    "'Verified' is set by code only if the URL appeared in that run's actual search results",
    "If web search ever fails, the researcher says so, removes links and marks everything 'unverified'",
    "Neither the AI nor the client can forge a source or its status",
], 13)

# 15 Revision
s = new_slide("Feedback in plain words → a revised model", "Step 4 · Logic model", "Step 4", BLUE)
picture(s, "13_logic_model_revised.jpg", 0.6, 1.5, w=7.0)
callout(s, 8.0, 1.5, 4.8, 2.25, "What the director asked", lm_feedback, YELLOW, 11.5)
callout(s, 8.0, 3.95, 4.8, 2.75, "What the team did (version 2)", lm_revision_msg, GREEN, 11.5)

# 16 Metrics research + review
s = new_slide("Designing metrics: research first, then peer review", "Step 5 · Metrics", "Step 5", GREEN)
picture(s, "14_metrics_research_live.jpg", 0.6, 1.55, w=6.0)
picture(s, "15_metrics_peer_review.jpg", 6.9, 1.55, w=6.0)
text(s, 0.6, 5.65, 6, 0.6, "The Funder Researcher searches for measures and real question items funders use.", 11, False, MUTED)
text(s, 6.9, 5.65, 6, 0.6, "Then all four specialists review the working plan in parallel.", 11, False, MUTED)

# 17 Plan
s = new_slide("The measurement plan: five metrics, two ways to collect", "Step 5 · Metrics", "Step 5", GREEN)
rows = [["Metric", "Method", "Questions on", "Outcome it tracks"]]
for m in plan["metrics"]:
    on = sum(q["included"] for q in m["questions"])
    rows.append([m["name"], m["method"].replace("_", " "), f"{on} of {len(m['questions'])}", m["outcome"]])
table(s, 0.6, 1.5, 12.1, rows, [3.6, 1.8, 1.3, 5.4], 11)
callout(s, 0.6, 4.25, 5.9, 2.4, "Director's feedback", metrics_feedback, YELLOW, 11.5)
callout(s, 6.8, 4.25, 5.9, 2.4, "Version 2", metrics_revision_msg[:420] + ("…" if len(metrics_revision_msg) > 420 else ""), GREEN, 11)

# 18 Question bank
s = new_slide("Several cited questions per metric, and you choose", "Step 5 · Metrics", "Step 5", GREEN)
picture(s, "17_metric_question_bank.jpg", 0.6, 1.5, w=7.0)
bullets(s, 8.0, 1.5, 4.8, 5.3, [
    "Each metric offers 3–5 candidate questions; the team switches on the 2–3 it recommends",
    "\"Why funders care\" links each metric to specific funders, with citations",
    "Good catch from the team: the stress items are inspired by the PSS-10 but are not copies, so it warned "
    "not to score them as a PSS-10 and to check licensing",
    f"{len(plan['sources'])} cited sources on the final plan, all verified",
], 12.5)

# 19 Done
s = new_slide("Approved, saved to the account, downloadable", "Step 5 · Done", "Step 5", GREEN)
picture(s, "18_studio_done.jpg", 0.6, 1.5, w=7.0)
bullets(s, 8.0, 1.5, 4.8, 5.3, [
    "Approving copies the logic model and plan onto the organization",
    "The team's working notes (research, brief, question pool) are deleted",
    "The PowerPoint download includes the logic model, plan, tool question lists, a funder alignment slide and sources",
    "Next: build the survey tools",
], 13)

# 20 Tools
s = new_slide("Survey tools are built straight from the metrics", "Step 6 · Collect", "Step 6", PURPLE)
picture(s, "19_tools_list.jpg", 0.6, 1.5, w=6.2)
bullets(s, 7.2, 1.5, 5.6, 5.3, [
    "One click turned the approved plan into tools:",
    "'Brief participant check-in': the 4 switched-on participant questions",
    "8 program numbers from the two 'program records' metrics (move-in and 12-month housing status), "
    "entered by staff and never sent as a survey",
    "Staff can hide questions, change timing, add open questions, mark reverse-scoring and pick demographics",
    "Once answers exist, wording and scales lock so pre and post stay comparable",
], 12.5)

# 21 Review bot
s = new_slide("The Survey Review Bot caught a real problem", "Step 6 · Collect", "Step 6", PURPLE)
picture(s, "21_review_bot_CROPTOP.jpg", 0.6, 1.5, w=7.0, crop_top=0.27)
callout(s, 8.0, 1.5, 4.8, 2.1, "Its summary (live)", review["summary"], PURPLE, 11)
bullets(s, 8.0, 3.8, 4.8, 3.0, [f"Applied: {a['action'].replace('_', ' ')}: {a['problem'][:85]}…" for a in review["applied"]]
        + ["Staff choose: fixes only change the draft until you Save"], 11.5)

# 22 Portal
s = new_slide("Participants get their own portal: a code and PIN, no names", "Step 6 · Collect", "Step 6", BLUE)
picture(s, "25_portal_list.jpg", 0.9, 1.45, h=5.35)
picture(s, "26_portal_form_EARLIER_TEST.jpg", 4.6, 1.45, h=5.35)
bullets(s, 8.2, 1.5, 4.6, 5.3, [
    "Staff create codes + 6-digit PINs, shown once (CSV/print); stored only as hashes",
    "Participants see only their own pending surveys and can submit pre and post once each",
    "Phone-first: big tap targets, optional demographics with 'Prefer not to say'",
    f"This run: {collect['participant_pre']} pre and {collect['participant_post']} post surveys from 24 participants",
    "(The form screenshot is from the earlier scripted test, same screen.)",
], 12)

# 23 Program numbers
s = new_slide("Numbers the organization already has: entered by staff", "Step 6 · Collect", "Step 6", PURPLE)
picture(s, "22_program_numbers.jpg", 0.6, 1.5, w=7.0)
bullets(s, 8.0, 1.5, 4.8, 5.3, [
    "No survey needed for things like move-ins, follow-ups, dollars spent or volunteers' work",
    "8 numbers came from the plan; 3 custom ones were added (rental assistance $, apartments furnished, landlord partners)",
    "Entered by month; totals and history shown; each appears as a dashboard chart",
], 13)

# 24 Dashboard
s = new_slide("The dashboard fills itself as data comes in", "Step 7 · Dashboard", "Step 7", YELLOW)
picture(s, "23_dashboard_top.jpg", 0.6, 1.5, w=7.0)
bullets(s, 8.0, 1.5, 4.8, 5.3, [
    f"{h['participants']} participants · {h['responses_pre']} pre / {h['responses_post']} post · {h['matched']} matched",
    f"Top result: {h['top_improvement']['pct']}% improved on {h['top_improvement']['metric'].lower()}",
    "Change is measured only for people who answered both pre and post",
    "Every chart has a table view and its own Excel download",
], 13)

# 25 Results (native charts)
s = new_slide("Results from the synthetic data (native, editable charts)", "Step 7 · Dashboard", "Step 7", YELLOW)
imp = chart_by_title("Share of participants who improved")
bar_chart(s, 0.6, 1.45, 6.1, 2.6, "Share of participants who improved", imp["categories"],
          [("% improved", imp["series"][0]["values"])], [BLUE], "%", max_v=100)
for i, prefix in enumerate(["Confidence addressing", "Knowing where", "Parent-reported stress"]):
    c = chart_by_title(prefix)
    title = c["title"].split(":")[0]
    bar_chart(s, 0.6 + (i % 2) * 6.2 if i < 2 else 6.8, 4.2 if i < 2 else 1.45, 6.0, 2.6,
              {"Confidence addressing a rent or landlord issue": "Confidence with a rent/landlord issue (0–10)",
               "Knowing where to seek help for a chosen need": "Knows where to seek help (% yes)",
               "Parent-reported stress": "Stress, reverse-scored (1–5, higher = less stress)"}.get(title, title),
              c["categories"], [(c["series"][0]["name"], c["series"][0]["values"])],
              [BLUE], c["unit"] if c["unit"] == "%" else "", max_v=100 if c["unit"] == "%" else None)
text(s, 0.6, 6.8, 12, 0.3, "Matched participants only. The stress items are reverse-scored, so a higher average means less stress.",
     10, False, MUTED)

# 26 Who we serve + program activity
s = new_slide("Who we serve, with privacy built in, and program activity", "Step 7 · Dashboard", "Step 7", YELLOW)
picture(s, "24_dashboard_demographics_CROP.jpg", 0.6, 1.5, w=4.6, crop_top=0.77, crop_bottom=0.0)
text(s, 5.35, 1.5, 1.15, 2.4, "Groups under 5 people are combined, so no one can be identified.", 10.5, False, MUTED)
ra = chart_by_title("Rental assistance paid")
line_chart(s, 6.6, 1.45, 6.1, 2.6, "Rental assistance paid to landlords ($)", ra["categories"],
           [("Rental assistance", ra["series"][0]["values"])], [BLUE])
due = chart_by_title("How many families are due")
housed = chart_by_title("How many families with a known status were in permanent")
bar_chart(s, 0.6, 4.35, 12.1, 2.4, "12-month follow-up: families due vs. in permanent housing", due["categories"],
          [("Due for 12-month follow-up", due["series"][0]["values"]),
           ("In permanent housing at follow-up", housed["series"][0]["values"])], [BLUE, RED], horizontal=False)

# 27 Downloads
s = new_slide("Everything downloads, in formats nonprofits already use", "Step 7 · Downloads", "Step 7", YELLOW)
bullets(s, 0.6, 1.5, 6.0, 5.2, [
    [("Each chart as Excel: ", True, INK), ("its data plus a native Excel chart", False, INK)],
    [("All data as Excel: ", True, INK), ("codebook, one sheet per tool (participant code × wave), program numbers", False, INK)],
    [("Impact report as PowerPoint: ", True, INK), ("theory of change, headline numbers, every chart (editable), quotes, method", False, INK)],
    [("Studio PowerPoint: ", True, INK), ("logic model, measurement plan, tool question lists, funder alignment, sources", False, INK)],
    "Text that starts with = + - @ is escaped so it can't run as a spreadsheet formula",
], 13.5)
callout(s, 7.0, 1.5, 5.7, 2.6, "From this run (in runs/2026-10-20_live_demo/)",
        "impact_report.pptx · survey_data.xlsx · dashboard.json · studio.json (full transcript, logic model, plan, sources) "
        "· run_log_tools_collect.json · participant_codes.json", GREEN, 12)
quotes = [q["text"] for q in dash["open_responses"][:3]]
callout(s, 7.0, 4.35, 5.7, 2.35, "\"In their words\" (synthetic reflections)", "\n".join(f"“{q}”" for q in quotes), RED, 11.5)

# 28 Under the hood
s = new_slide("Under the hood: cost, time and accountability", "Behind the scenes", accent=INK)
tiles = [(f"{TOKENS['runs']}", "agent runs"), (f"{TOKENS['requests']}", "model requests"),
         (f"{TOKENS['input'] // 1000}K / {TOKENS['output'] // 1000}K", "input / output tokens"),
         ("~18 min", "whole live test"), (f"{TOKENS['refused']}", "extra consults refused by budget"),
         (f"{TOKENS['audit']}", "audit entries, hash chain intact")]
for i, (big, small) in enumerate(tiles):
    x, y = 0.6 + (i % 3) * 4.1, 1.55 + (i // 3) * 1.65
    shape(s, MSO_SHAPE.RECTANGLE, x, y, 3.85, 1.4, WHITE, LINE)
    text(s, x + 0.2, y + 0.15, 3.5, 0.6, big, 26, True, [BLUE, RED, GREEN, YELLOW, PURPLE, INK][i], HEAD)
    text(s, x + 0.2, y + 0.82, 3.5, 0.5, small, 12, False, MUTED)
bullets(s, 0.6, 5.05, 12.1, 1.8, [
    "Most tokens went to the Team Lead (≈155K in) and the Funder Researcher's web research (≈81K in); the Survey Review Bot used ≈4K",
    "Every agent step, hand-off, approval, download and sign-in is in output/audit_log.json; no passwords, PINs or uploaded text",
], 12.5)

# 29 Findings
s = new_slide("What the test taught us, and what I fixed on the spot", "Findings", accent=RED)
bullets(s, 0.6, 1.5, 6.0, 5.3, [
    [("Working well: ", True, GREEN), ("live web research with verified local funders; real peer review; metrics that respect licensing and privacy", False, INK)],
    [("Working well: ", True, GREEN), ("the Survey Review Bot caught reverse-scoring and a missing open question", False, INK)],
    [("Working well: ", True, GREEN), ("pre/post matching, suppression of small groups, all downloads", False, INK)],
    [("To watch: ", True, YELLOW), ("intake is fast (~15 s); drafting and metrics take 1–4 minutes each", False, INK)],
], 13)
bullets(s, 6.9, 1.5, 5.8, 5.3, [
    [("Fixed: ", True, BLUE), ("test answer matching now uses whole words (it mis-sent one answer)", False, INK)],
    [("Fixed: ", True, BLUE), ("program numbers matched to the AI's indicator names, with realistic values; re-runs replace, not duplicate", False, INK)],
    [("Fixed: ", True, BLUE), ("charts now say 'higher = better' for reverse-scored metrics", False, INK)],
    [("Fixed: ", True, BLUE), ("squished 'Send to team' button; long tool descriptions clamped", False, INK)],
    [("Next: ", True, RED), ("Privacy Guard for quotes, password reset, PDF report", False, INK)],
], 13)

# 30 Close
s = new_slide("Try it yourself", "Next", accent=GREEN)
bullets(s, 0.6, 1.6, 12, 4.5, [
    f"Sign in as {profile['account']['username']} (password in profile.json) to explore Keystone's real results",
    "Participant portal: /participant, using any code + PIN from participant_codes.json",
    "Re-run everything: .venv/Scripts/python scripts/run_test_profile.py data/test_profiles/keystone_family_housing",
    "Rebuild this deck: .venv/Scripts/python scripts/build_demo_deck.py data/test_profiles/keystone_family_housing 2026-10-20_live_demo",
], 15)
marks(s, 0.6, 6.2, 0.35)

out = run_dir / "Through_Line_live_walkthrough.pptx"
prs.save(out)
print(f"saved {out} ({len(prs.slides)} slides)")
