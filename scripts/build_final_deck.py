"""Build the final-presentation deck (docs/Through_Line_final_deck.pptx) in Through Line's branding.

    .venv/Scripts/python scripts/build_final_deck.py

Same story and numbers as the report (docs/report/index.html); images come from docs/report/img/.
Values marked PENDING are filled in once the live-site usage and tester feedback are in.
"""

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "docs" / "report" / "img"
OUT = ROOT / "docs" / "Through_Line_final_deck.pptx"

RED, BLUE, YELLOW, GREEN, PURPLE = "D96C4F", "3E63B0", "E8B84A", "4E8C7A", "8A6FB0"
INK, MUTED, PAPER, CARD, LINE = "2B2B2B", "6F6A63", "F7F4EE", "FFFFFF", "E5DFD5"
TINT = {"red": "FBEFEA", "blue": "EDF1F9", "yellow": "FCF5E4", "green": "EBF3F0", "purple": "F2EEF8"}
HEAD, BODY = "Century Gothic", "Segoe UI"


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def shape(s, kind, x, y, w, h, fill, line=None):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, value, size=16, bold=False, color=INK, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         italic=False, space=4):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.02)
    lines = value if isinstance(value, list) else [value]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        runs = line if isinstance(line, list) else [(line, {})]
        for part in runs:
            t, fmt = (part, {}) if isinstance(part, str) else part
            r = p.add_run()
            r.text = t
            f = r.font
            f.size, f.bold, f.italic, f.name = Pt(fmt.get("size", size)), fmt.get("bold", bold), italic, fmt.get("font", font)
            f.color.rgb = rgb(fmt.get("color", color))
    return tb


def marks(s, x, y, size=0.22):
    shape(s, MSO_SHAPE.OVAL, x, y, size, size, RED)
    shape(s, MSO_SHAPE.RECTANGLE, x + size * 1.25, y, size, size, BLUE)
    shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, x + size * 2.5, y, size * 1.1, size, YELLOW)


def page(title: str, eyebrow: str, n: int, accent: str, notes: str):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(PAPER)
    shape(s, MSO_SHAPE.RECTANGLE, 0.6, 0.55, 0.55, 0.07, accent)
    text(s, 0.6, 0.7, 10, 0.35, eyebrow.upper(), 11, True, MUTED)
    text(s, 0.6, 1.0, 12.1, 1.0, title, 30, True, INK, HEAD)
    marks(s, 0.6, 7.0, 0.13)
    text(s, 1.15, 6.92, 6, 0.3, "Through Line · AI Foundations · Yale SOM", 10, False, MUTED)
    text(s, 12.2, 6.92, 0.6, 0.3, str(n), 10, True, MUTED, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


def card(s, x, y, w, h, fill=CARD, edge=None):
    c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, LINE)
    c.adjustments[0] = 0.06
    if edge:
        shape(s, MSO_SHAPE.RECTANGLE, x, y, w, 0.08, edge)
    return c


def stat(s, x, y, w, h, value, label, edge, value_size=40):
    card(s, x, y, w, h, CARD, edge)
    text(s, x + 0.3, y + 0.3, w - 0.6, 1.0, value, value_size, True, INK, HEAD)
    text(s, x + 0.3, y + 1.35, w - 0.6, h - 1.5, label, 14, False, INK)


def picture(s, name, x, y, w=None, h=None):
    path = IMG / name
    return s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def pending(s, x, y, w, h, msg):
    box = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, TINT["yellow"], "C99A2E")
    box.line.dash_style = 4
    text(s, x + 0.2, y + 0.12, w - 0.4, h - 0.2, [[("PENDING  ", {"bold": True, "color": "7A5A12"}), (msg, {"color": "5B4510"})]], 13)


# ============================================================ 1. title
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = rgb(PAPER)
marks(s, 0.8, 0.8, 0.26)
text(s, 1.85, 0.72, 4, 0.5, "Through Line", 20, True, INK, HEAD)
text(s, 0.8, 1.75, 7.2, 2.6, "From mission to evidence to story.", 54, True, INK, HEAD, space=0)
text(s, 0.8, 4.35, 6.8, 1.2, "AI agents that help small nonprofits define their impact, measure it, and turn it into funder "
     "reports and grant applications, without an evaluation consultant.", 18, False, MUTED)
text(s, 0.8, 5.75, 7.5, 0.9, [[("Live  ", {"bold": True}), "through-line.onrender.com"],
                              [("Code  ", {"bold": True}), "github.com/doug-palmer28/through-line"]], 14, False, INK)
text(s, 0.8, 6.75, 7, 0.4, "Doug Palmer · AI Foundations final project · Yale School of Management", 11, False, MUTED)
shape(s, MSO_SHAPE.RECTANGLE, 8.2, 4.45, 4.6, 0.05, INK)
shape(s, MSO_SHAPE.OVAL, 9.3, 1.3, 2.4, 2.4, RED).fill.fore_color.brightness = 0
shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, 8.4, 3.4, 2.4, 2.4, YELLOW)
shape(s, MSO_SHAPE.RECTANGLE, 10.9, 3.9, 1.9, 1.9, BLUE)
s.notes_slide.notes_text_frame.text = (
    "Through Line puts a program-evaluation team in a small nonprofit's hands: a team of AI agents helps staff build a logic "
    "model, choose metrics, run pre/post surveys, and turn the results into funder decks and grant drafts.")

# ============================================================ 2. the problem
s = page("Funders want proof. Small nonprofits can't afford it.", "The problem and who has it", 2, RED,
         "Funders increasingly want outcome data. Most nonprofits have no evaluation staff or budget, and the smallest get "
         "almost none of the foundation money. That creates a cycle where the organizations that can pay for evaluation "
         "get more funding. Our user: the executive director of a one-program nonprofit, writing grant reports at night.")
stat(s, 0.6, 2.3, 3.9, 2.7, "8%", "of nonprofits have evaluation staff. In 63%, the director or program staff do it themselves.", RED)
stat(s, 4.72, 2.3, 3.9, 2.7, "84%", "spend under 5% of their budget on evaluation, below the 5–10% evaluators recommend.", YELLOW)
stat(s, 8.84, 2.3, 3.9, 2.7, "0.4%", "of foundation funding went to the ~60% of nonprofits with budgets under $50K (2019–2023).", BLUE)
card(s, 0.6, 5.25, 12.14, 1.35, TINT["red"])
text(s, 0.9, 5.38, 11.6, 1.2, [[("Who has it: ", {"bold": True}), "the executive director of a small, one-program nonprofit, often "
     "serving those furthest from opportunity. "], [("Inequity compounds: ", {"bold": True}),
     "Black-led organizations hold 76% smaller unrestricted net assets than white-led peers."]], 15, False, INK)
text(s, 0.6, 6.62, 12, 0.3, "Sources: Innovation Network, State of Evaluation (2016); Candid (2025); Echoing Green & Bridgespan (2020).",
     9, False, MUTED)

# ============================================================ 3. what's new
s = page("An evaluation team in the nonprofit's own hands", "What's new", 3, BLUE,
         "Today this is bought from consultants, at $3,000 or more for a logic model alone and $150–250 an hour after that. "
         "Software tools each cover one piece. Through Line is new in two ways: the nonprofit does it itself, guided by an "
         "AI team, and everything flows together, from logic model to grant draft, with live numbers.")
card(s, 0.6, 2.25, 5.6, 3.55, CARD)
text(s, 0.9, 2.45, 5.1, 0.4, "TODAY", 12, True, MUTED)
text(s, 0.9, 2.85, 5.1, 2.9, [[("$3,000–3,500  ", {"bold": True, "font": HEAD, "size": 20}), "a consultant-led logic model"],
                              [("$150–250/hr  ", {"bold": True, "font": HEAD, "size": 20}), "evaluators for plans and surveys"],
                              [("$299–999/mo  ", {"bold": True, "font": HEAD, "size": 20}), "software that covers one piece "
                               "(Sopact, Instrumentl)"],
                              ["Nothing connects the logic model to the data or the grant."]], 15, False, INK, space=10)
card(s, 6.5, 2.25, 6.24, 3.55, TINT["blue"])
text(s, 6.8, 2.45, 5.7, 0.4, "THROUGH LINE", 12, True, BLUE)
text(s, 6.8, 2.85, 5.7, 2.9, [[("1. The nonprofit does it itself. ", {"bold": True}),
                               "Five AI specialists play the consultant team, with verified sources. Staff answer, edit, approve."],
                              [("2. Everything flows together. ", {"bold": True}),
                               "Outcomes become metrics, which become surveys, which fill the dashboard, deck, and grant draft."],
                              [("Cost: ", {"bold": True}), "cents per organization in AI fees."]], 15, False, INK, space=10)
steps = [("Logic model", RED), ("Metrics", "C9603F"), ("Surveys", BLUE), ("Data", "3456A0"), ("Dashboard", YELLOW),
         ("Funder deck", GREEN), ("Grant draft", PURPLE)]
w = 12.14 / len(steps)
for i, (label, color) in enumerate(steps):
    sh = shape(s, MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, 0.6 + i * w, 6.0, w + 0.05, 0.6, color)
    text(s, 0.6 + i * w + 0.1, 6.0, w - 0.1, 0.6, label, 13, True, INK if color == YELLOW else "FFFFFF", HEAD,
         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# ============================================================ 4. agents
s = page("Nine agents, one model, plain-code orchestration", "Agents", 4, YELLOW,
         "The Studio is a team of five: the Team Lead talks to staff, and four specialists propose questions and peer-review "
         "every draft. The Funder Researcher uses live web search, and code verifies its citations. The Survey Review Bot, "
         "Report Designer, Grant Writer, and Privacy Reviewer each handle one feature. Everything runs on gpt-6-luna for cost, "
         "and it held up in testing. The orchestrator is plain Python, so limits and order can't be talked around.")
picture(s, "agents_diagram.png", 0.6, 1.95, w=8.6)
card(s, 9.45, 1.95, 3.29, 4.75, CARD, YELLOW)
text(s, 9.7, 2.2, 2.85, 4.5, [[("Model", {"bold": True, "font": HEAD, "size": 17})],
                              "gpt-6-luna for every agent, via Portkey (Responses API). $0.10 / $0.50 per million tokens: "
                              "20× cheaper than gpt-6-sol, and good enough in testing.",
                              [("Tools", {"bold": True, "font": HEAD, "size": 17})],
                              "Live web search (Funder Researcher), consult and context tools (Studio). Everyone else "
                              "works only from the data code gives them.",
                              [("Typed outputs", {"bold": True, "font": HEAD, "size": 17})],
                              "Every agent returns a checked Pydantic object."], 13, False, INK, space=6)

# ============================================================ 5. demo
s = page("Keystone Family Housing, from documents to grant draft", "Demo · a live run on gpt-6-luna", 5, INK,
         "A real run with a fictional rapid-rehousing nonprofit:\n"
         "- 10.5 minutes from uploaded documents to an approved logic model and five metrics with 20 verified sources.\n"
         "- 24 participants answered the surveys through the portal.\n"
         "- The dashboard filled itself in.\n"
         "- The Report Designer built a 14-slide deck in about 2 minutes.\n"
         "- The grant draft honestly flagged that Keystone may not be eligible for a Colorado-only foundation.")
grid = [("live_07_interview_thinking.jpg", "Five questions, with the live team huddle"),
        ("live_10_logic_model_grid.jpg", "Logic model in their own words"),
        ("live_16_metrics_plan.jpg", "Metrics with cited questions"),
        ("dashboard.jpg", "Dashboard: 85% · 60% · 70% improved"),
        ("preview_funder_deck.jpg", "14-slide funder deck, live numbers"),
        ("grants_answers.jpg", "Grant draft for one foundation")]
for i, (name, cap) in enumerate(grid):
    x, y = 0.6 + (i % 3) * 4.1, 1.95 + (i // 3) * 2.45
    card(s, x, y, 3.95, 2.3, CARD)
    pic = picture(s, name, x + 0.1, y + 0.1, w=3.75)
    if pic.height > Inches(1.75):  # crop tall images to the frame
        ratio = Inches(1.75) / pic.height
        pic.crop_bottom = 1 - ratio
        pic.height = Inches(1.75)
    text(s, x + 0.15, y + 1.9, 3.7, 0.35, cap, 12, True, INK)

# ============================================================ 6. cost
s = page("Cents per organization, not thousands of dollars", "Cost to run", 6, GREEN,
         "Keystone's whole journey cost 6 cents in model fees: 33 agent runs, 339,000 input and 52,000 output tokens. Web search adds "
         "at most 18 cents. A typical year for one organization is about $1, and the most anyone can spend while staying "
         "within the allowances is about $8. That's $10,000 a year in AI fees for 10,000 organizations. The main cost control "
         "is limiting how often each organization can use the agents.")
stat(s, 0.6, 2.05, 3.6, 2.05, "$0.24", "Keystone's whole journey, upper bound ($0.06 tokens + ≤ $0.18 web search)", GREEN, 36)
cd = CategoryChartData()
cd.categories = ["Logic Model Studio", "Grant draft", "Impact report", "Survey review"]
cd.add_series("Model tokens (¢)", (4.93, 0.60, 0.35, 0.10))
cd.add_series("Web search, max (¢)", (12, 6, 0, 0))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(4.45), Inches(1.95), Inches(4.6), Inches(2.3), cd)
ch = gf.chart
ch.has_legend = True
ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
ch.legend.font.size, ch.legend.font.name = Pt(10), BODY
ch.category_axis.reverse_order = True
ch.category_axis.tick_labels.font.size = Pt(10)
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
for i, ser in enumerate(ch.plots[0].series):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = rgb(GREEN if i == 0 else "C8D7D1")
ch.plots[0].gap_width = 60
text(s, 4.45, 1.75, 4.6, 0.3, "Keystone cost by task, in cents", 11, True, MUTED)
card(s, 9.3, 2.05, 3.44, 2.05, CARD)
text(s, 9.5, 2.15, 3.1, 1.9, [[("At scale (AI fees / year)", {"bold": True})],
                              [("100 orgs  ", {"bold": True, "font": HEAD, "size": 18}), "≈ $100"],
                              [("10,000 orgs  ", {"bold": True, "font": HEAD, "size": 18}), "≈ $10,000"],
                              [("Worst case ", {"bold": True}), "≈ $8 per org"]], 13, False, INK, space=4)
card(s, 0.6, 4.45, 12.14, 2.2, TINT["green"])
text(s, 0.9, 4.6, 11.6, 2.0, [[("Keeping costs down: limit how often the bots run", {"bold": True, "font": HEAD, "size": 16})],
                              "Studio: 2 sessions, 5 questions, ≤ 2 revision rounds   ·   Survey Review Bot: 2 per survey   ·   "
                              "Impact report: 1 design + 1 feedback round, then free refreshes   ·   Grants: 3 a month   ·   "
                              "Privacy check: any time (unchanged drafts are free)",
                              [("Did it make the app worse? ", {"bold": True}), "In our runs, rarely: Keystone used 1 of 2 revision "
                               "rounds, and refused consults didn't hurt quality. Testers' view: PENDING."]], 13, False, INK, space=6)

# ============================================================ 7. safety
s = page("Built for people's private information", "Safety", 7, RED,
         "Four layers of protection:\n"
         "- Staff accounts: Argon2id passwords, usernames stored only as hashes, strict cookies, lockouts, and organizations walled off from each other.\n"
         "- Participants are codes and PINs, never names.\n"
         "- A Privacy Reviewer agent plus code scans check everything that leaves the organization.\n"
         "- Agents read only what a person can see: our test PDF hides 'ignore your instructions' six ways, and the standard reader "
         "pulls out all six while ours pulls out none.\n"
         "Usage limits stop overspending.")
pillars = [("Staff accounts", "Argon2id passwords · usernames stored only as hashes · strict cookies · lockouts · each organization walled off", BLUE),
           ("Codes, not names", "Participants sign in with a code + PIN. No names or contact details, ever. Groups under 5 hidden.", YELLOW),
           ("Privacy Reviewer", "Code scrub → AI review → code scan on every deck, flyer, and grant draft before it leaves", RED),
           ("Usage limits", "Per-organization allowances + per-run token caps stop overspending", GREEN)]
for i, (head, body, color) in enumerate(pillars):
    y = 1.95 + i * 1.18
    card(s, 0.6, y, 6.6, 1.05, CARD)
    shape(s, MSO_SHAPE.RECTANGLE, 0.6, y, 0.1, 1.05, color)
    text(s, 0.9, y + 0.1, 6.2, 0.4, head, 16, True, INK, HEAD)
    text(s, 0.9, y + 0.48, 6.2, 0.55, body, 12, False, INK)
card(s, 7.45, 1.95, 5.29, 4.6, CARD, PURPLE)
text(s, 7.7, 2.15, 4.8, 0.5, "Agents read only what a person can see", 16, True, INK, HEAD)
picture(s, "hidden_text_page.jpg", 7.75, 2.75, w=4.7)
text(s, 7.7, 5.5, 4.85, 1.0, [[("Hidden 6 ways: ", {"bold": True}), "white, tiny, off-page, invisible, transparent, under a box."],
                              [("Standard reader: ", {"bold": True}), "extracts all 6.  ", ("Through Line: ", {"bold": True}), "0."]],
     12, False, INK, space=2)

# ============================================================ 8. impact
s = page("What it does for an organization", "Impact · the numbers", 8, PURPLE,
         "From Keystone's live run:\n"
         "- About 10 minutes to an approved logic model and metrics, with 20 verified sources.\n"
         "- Surveys that 24 participants actually completed, with 85% of matched participants improving on housing confidence.\n"
         "- A 14-slide funder deck in about 2 minutes.\n"
         "- All for under a dollar.\n"
         "Fill in the testers' numbers and quotes once they're in.")
nums = [("~10 min", "documents → approved logic model + 5 metrics", RED),
        ("20", "verified sources behind the metrics", BLUE),
        ("85%", "of matched participants improved on housing confidence", YELLOW),
        ("< $1", "AI cost, vs. $3,000+ for a consultant-led logic model", GREEN)]
for i, (v, lab, color) in enumerate(nums):
    stat(s, 0.6 + i * 3.08, 1.95, 2.93, 2.35, v, lab, color, 34)
pending(s, 0.6, 4.55, 6.0, 2.05, "Kids at Their Best (Jodi Walker): how far she got, time spent, would she use it, and one quote.")
pending(s, 6.74, 4.55, 6.0, 2.05, "Sara: how far she got, time spent, would she use it, and one quote.")

# ============================================================ 9. what's next
s = page("What's next", "Limits and next steps", 9, INK,
         "Limits:\n"
         "- It works best for a nonprofit with one core program.\n"
         "- Getting data in from other systems is the scaling challenge.\n"
         "Next steps:\n"
         "- Multi-program logic models.\n"
         "- An MCP server, so a staff member's AI assistant can move data from Google Sheets, Airtable, or Salesforce, which already "
         "have official MCP servers.\n"
         "- More chart types and report formats.\n"
         "- The qualitative team: privacy guard, outcome coder, testimonial writer.")
cols = [("Beyond one program", "Logic models per program, rolled up across the organization; shared participants and outcomes.", RED),
        ("Connect their data (MCP)", "A Through Line MCP server (import responses, add program numbers, read the dashboard) that "
         "works with the official Google Sheets, Airtable, and Salesforce MCP servers.", BLUE),
        ("Richer reporting", "More chart types, benchmarks from the field, annual-report pages, and coding of open answers.", YELLOW)]
for i, (head, body, color) in enumerate(cols):
    x = 0.6 + i * 4.1
    card(s, x, 2.0, 3.95, 3.2, CARD, color)
    text(s, x + 0.3, 2.3, 3.4, 0.9, head, 19, True, INK, HEAD)
    text(s, x + 0.3, 3.15, 3.4, 2.0, body, 14, False, INK)
card(s, 0.6, 5.45, 12.14, 1.15, TINT["yellow"])
text(s, 0.9, 5.58, 11.6, 1.0, [[("Fixes we'd make first: ", {"bold": True}), "catch near-duplicate interview questions; a library of "
     "validated, licensed survey scales; and check that a cited page actually supports the claim, not just that it exists."]],
     14, False, INK)

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
