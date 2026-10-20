"""Operator tools for whoever runs Through Line: token usage across every organization, and allowances.

    .venv/Scripts/python scripts/operator_tools.py usage               # print totals + write an Excel report
    .venv/Scripts/python scripts/operator_tools.py orgs                # organizations and what they've used
    .venv/Scripts/python scripts/operator_tools.py report-data         # compact JSON: tokens by task, what each org built
    .venv/Scripts/python scripts/operator_tools.py backfill            # once: copy pre-ledger usage from the audit log
    .venv/Scripts/python scripts/operator_tools.py export-org 1 scripts/demo/keystone.json   # package an organization
    python ../scripts/operator_tools.py import-org ../scripts/demo/keystone.json --username keystone.demo   # on Render
    .venv/Scripts/python scripts/operator_tools.py allow 1 grants      # give org 1 one more grant this month
    .venv/Scripts/python scripts/operator_tools.py allow 1 report-revision --item 3

Allowance kinds: studio, report, grants, report-revision (--item report id),
grant-revision (--item application id), survey-review (--item survey tool id).
"allow" lowers the used count by one (or --n); it never deletes anything, and it's audited.

Reads the real database (data/through_line.db) unless THROUGH_LINE_DB is set. The Excel
report goes to data/usage_reports/ (git-ignored): it names staff and organizations.
"""

import argparse
import os
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from sqlmodel import Session, select  # noqa: E402

import allowances  # noqa: E402
import ledger  # noqa: E402
from audit import audit_append  # noqa: E402
from db import GrantApplication, Instrument, Organization, ReportRun, engine  # noqa: E402
from reports import excel_safe  # noqa: E402

NUMS = ("requests", "input_tokens", "output_tokens", "total_tokens")


def _sum(rows, key) -> list[tuple]:
    groups: dict = defaultdict(lambda: [0, 0, 0, 0])
    for r in rows:
        g = groups[key(r)]
        for i, n in enumerate(NUMS):
            g[i] += r[n]
    return [(*(k if isinstance(k, tuple) else (k,)), *v) for k, v in sorted(groups.items())]


def _sheet(wb, title: str, header: list[str], rows) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="23262B")
    for row in rows:
        ws.append([excel_safe(v) for v in row])  # org and staff names are user-typed: never formulas
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, min(40, max(len(str(c.value or "")) for c in col) + 2))
    ws.freeze_panes = "A2"


def usage_report(out_dir: Path | None = None) -> Path:
    with Session(engine()) as db:
        rows = ledger.platform_rows(db)
    wb = Workbook()
    wb.remove(wb.active)
    totals = ["Requests", "Input tokens", "Output tokens", "Total tokens"]
    _sheet(wb, "By organization", ["Organization", *totals], _sum(rows, lambda r: r["organization"]))
    _sheet(wb, "By person", ["Organization", "Person", *totals], _sum(rows, lambda r: (r["organization"], r["person"])))
    _sheet(wb, "By feature", ["Organization", "Feature", *totals], _sum(rows, lambda r: (r["organization"], r["feature"])))
    _sheet(wb, "By month", ["Month", "Organization", *totals], _sum(rows, lambda r: (r["when"][:7], r["organization"])))
    _sheet(wb, "Every run", list(rows[0].keys()) if rows else ["(no runs yet)"], [list(r.values()) for r in rows])
    out_dir = out_dir or ROOT / "data" / "usage_reports"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"usage_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    wb.save(path)
    return path


def _feature_of(phase: str) -> str:
    if phase == "survey_review":
        return "survey_review"
    if phase.startswith("report") or phase == "privacy_check":
        return "impact_report"
    if phase.startswith("grant"):
        return "grants"
    return "studio"


def backfill() -> str:
    """Copy token use recorded in the audit log before the ledger existed (person unknown). Safe to re-run."""
    import json
    import os
    import re

    from config import AUDIT_LOG_PATH
    from db import TokenUsage

    log = json.loads(Path(os.getenv("THROUGH_LINE_AUDIT", AUDIT_LOG_PATH)).read_text(encoding="utf-8"))
    sessions_dir = ROOT / "data" / "sessions"
    org_of_session: dict[str, int | None] = {}

    def org_for(session_id: str | None) -> int | None:
        if not session_id:
            return None
        m = re.match(r"org-(\d+)-", session_id)
        if m:
            return int(m.group(1))
        if session_id not in org_of_session:
            path = sessions_dir / f"{session_id}.json"
            org_of_session[session_id] = json.loads(path.read_text(encoding="utf-8")).get("org_id") if path.exists() else None
        return org_of_session[session_id]

    added = skipped = 0
    with Session(engine()) as db:
        done = {r.run_id for r in db.exec(select(TokenUsage).where(TokenUsage.run_id.startswith("audit:"))).all()}
        first_live = min((r.at for r in db.exec(select(TokenUsage).where(~TokenUsage.run_id.startswith("audit:"))).all()),
                         default=None)
        for e in log:
            if e.get("event") not in ("agent_finished", "agent_failed", "limit_reached") or not e.get("usage"):
                continue
            key = f"audit:{e['seq']}"
            at = datetime.fromisoformat(e["logged_at"])
            if key in done or (first_live and at.replace(tzinfo=None) >= first_live.replace(tzinfo=None)):
                continue  # already copied, or already recorded live by the ledger
            org_id = org_for(e.get("session_id"))
            if org_id is None:
                skipped += 1
                continue
            u = e["usage"]
            db.add(TokenUsage(at=at, org_id=org_id, user_id=None, feature=_feature_of(e.get("phase") or ""),
                              phase=e.get("phase") or "", agent=e.get("agent") or "", run_id=key,
                              requests=u.get("requests", 0), input_tokens=u.get("input_tokens", 0),
                              output_tokens=u.get("output_tokens", 0), ok=e["event"] == "agent_finished"))
            added += 1
        db.commit()
    return f"Copied {added} earlier agent runs from the audit log ({skipped} had no organization, e.g. pre-accounts tests)."


# ============================================================ copying an organization (e.g. the Keystone demo)
#
# export-org writes one JSON file with an organization's work: logic model and plan, survey tools, participants
# (codes only), responses, program numbers, brand + logo, reports, grant drafts, and its Studio record.
# It never includes passwords, PINs, usernames, or token usage. import-org loads that file into another
# database (e.g. the live site) under a NEW staff login you choose; participants get fresh random PINs.


def _iso(dt) -> str | None:
    return dt.isoformat() if dt else None


def export_org(org_id: int, out: Path) -> str:
    import base64
    import json

    from branding import load_logo
    from db import Indicator, IndicatorEntry, Participant, Response
    from sessions import sessions_dir

    with Session(engine()) as db:
        org = db.get(Organization, org_id)
        if not org:
            raise SystemExit(f"No organization {org_id}.")
        by = lambda model: db.exec(select(model).where(model.org_id == org_id)).all()  # noqa: E731
        people = {p.id: p.code for p in by(Participant)}
        session_path = sessions_dir() / f"{org.studio_session_id}.json"
        logo = load_logo(org_id)
        bundle = {
            "format": "through-line-org/1", "exported_at": datetime.now().isoformat(timespec="seconds"),
            "organization": {"name": org.name, "logic_model": org.logic_model, "measurement_plan": org.measurement_plan,
                             "brand": org.brand},
            "logo_png_base64": base64.b64encode(logo).decode() if logo else None,
            "studio_session": json.loads(session_path.read_text(encoding="utf-8")) if session_path.exists() else None,
            "participants": [{"code": code} for code in people.values()],
            "instruments": [{"old_id": i.id, "name": i.name, "method": i.method, "when": i.when, "status": i.status,
                             "questions": i.questions, "demographics": i.demographics, "review": i.review,
                             "review_count": i.review_count, "created_at": _iso(i.created_at)} for i in by(Instrument)],
            "responses": [{"instrument": r.instrument_id, "participant": people.get(r.participant_id), "wave": r.wave,
                           "answers": r.answers, "demographics": r.demographics, "entered_by": r.entered_by,
                           "created_at": _iso(r.created_at)} for r in by(Response)],
            "indicators": [{"old_id": i.id, "name": i.name, "unit": i.unit, "description": i.description,
                            "source": i.source} for i in by(Indicator)],
            "entries": [{"indicator": e.indicator_id, "period": e.period.isoformat(), "value": e.value, "note": e.note}
                        for e in by(IndicatorEntry)],
            "reports": [{"request": r.request, "draft": r.draft, "privacy": r.privacy, "revisions": r.revisions,
                         "feedback": r.feedback, "created_at": _iso(r.created_at)} for r in by(ReportRun)],
            "grants": [{"foundation": g.foundation, "foundation_url": g.foundation_url, "request": g.request,
                        "sources": g.sources, "profile": g.profile, "research_mode": g.research_mode, "draft": g.draft,
                        "privacy": g.privacy, "revisions": g.revisions, "feedback": g.feedback,
                        "created_at": _iso(g.created_at)} for g in by(GrantApplication)],
        }
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(bundle, indent=1, ensure_ascii=False), encoding="utf-8")
    return (f"Exported '{bundle['organization']['name']}' to {out}: {len(bundle['participants'])} participants, "
            f"{len(bundle['responses'])} responses, {len(bundle['entries'])} program values, "
            f"{len(bundle['reports'])} report(s), {len(bundle['grants'])} grant draft(s). No passwords or PINs.")


def import_org(path: Path, username: str, password: str, display_name: str) -> str:
    import base64
    import json
    import re
    import secrets
    from datetime import date

    from branding import save_logo
    from db import Indicator, IndicatorEntry, Participant, Response, StaffUser
    from security import hash_secret, username_hash
    from sessions import sessions_dir

    b = json.loads(path.read_text(encoding="utf-8"))
    if b.get("format") != "through-line-org/1":
        raise SystemExit("That file isn't a Through Line organization export.")
    when = lambda s: datetime.fromisoformat(s) if s else datetime.now()  # noqa: E731
    with Session(engine()) as db:
        if db.exec(select(StaffUser).where(StaffUser.username_hash == username_hash(username))).first():
            raise SystemExit(f"The username '{username}' is taken. Choose another with --username.")
        o = b["organization"]
        org = Organization(name=o["name"], logic_model=o["logic_model"], measurement_plan=o["measurement_plan"],
                           brand=o["brand"], allowance_used={"studio_starts": 1, "report_designs": len(b["reports"])})
        db.add(org)
        db.commit()
        db.refresh(org)
        db.add(StaffUser(org_id=org.id, username_hash=username_hash(username), password_hash=hash_secret(password),
                         display_name=display_name, role="admin"))
        people = {}
        for p in b["participants"]:  # fresh random PINs: nobody (including the exporter) knows them
            row = Participant(org_id=org.id, code=p["code"], pin_hash=hash_secret(f"{secrets.randbelow(10**6):06d}"))
            db.add(row)
            db.flush()
            people[p["code"]] = row.id
        inst_ids, ind_ids = {}, {}
        for i in b["instruments"]:
            row = Instrument(org_id=org.id, name=i["name"], method=i["method"], when=i["when"], status=i["status"],
                             questions=i["questions"], demographics=i["demographics"], review=i["review"],
                             review_count=i.get("review_count") or 0, created_at=when(i.get("created_at")))
            db.add(row)
            db.flush()
            inst_ids[i["old_id"]] = row.id
        for r in b["responses"]:
            if r["participant"] in people and r["instrument"] in inst_ids:
                db.add(Response(org_id=org.id, instrument_id=inst_ids[r["instrument"]], participant_id=people[r["participant"]],
                                wave=r["wave"], answers=r["answers"], demographics=r["demographics"],
                                entered_by=r["entered_by"], created_at=when(r.get("created_at"))))
        for i in b["indicators"]:
            row = Indicator(org_id=org.id, name=i["name"], unit=i["unit"], description=i["description"], source=i["source"])
            db.add(row)
            db.flush()
            ind_ids[i["old_id"]] = row.id
        for e in b["entries"]:
            if e["indicator"] in ind_ids:
                db.add(IndicatorEntry(org_id=org.id, indicator_id=ind_ids[e["indicator"]],
                                      period=date.fromisoformat(e["period"]), value=e["value"], note=e["note"]))

        def remap(obj):
            """Chart ids embed database ids (metric-<tool>-…, choice-<tool>-…, indicator-<n>): point them at the new rows."""
            text = json.dumps(obj)
            text = re.sub(r'"(metric|choice)-(\d+)-', lambda m: f'"{m.group(1)}-{inst_ids.get(int(m.group(2)), m.group(2))}-', text)
            text = re.sub(r'"indicator-(\d+)"', lambda m: f'"indicator-{ind_ids.get(int(m.group(1)), m.group(1))}"', text)
            return json.loads(text)

        for r in b["reports"]:
            db.add(ReportRun(org_id=org.id, request=remap(r["request"]), draft=remap(r["draft"]), privacy=r["privacy"],
                             revisions=r.get("revisions") or 0, feedback=r.get("feedback") or [],
                             created_at=when(r.get("created_at"))))
        for g in b["grants"]:
            db.add(GrantApplication(org_id=org.id, foundation=g["foundation"], foundation_url=g["foundation_url"],
                                    request=g["request"], sources=g["sources"], profile=g["profile"],
                                    research_mode=g["research_mode"], draft=g["draft"], privacy=g["privacy"],
                                    revisions=g.get("revisions") or 0, feedback=g.get("feedback") or [],
                                    created_at=when(g.get("created_at"))))
        if b.get("studio_session"):  # the Studio record, re-owned by the new organization
            s = {**b["studio_session"], "org_id": org.id}
            sessions_dir().mkdir(parents=True, exist_ok=True)
            (sessions_dir() / f"{s['session_id']}.json").write_text(json.dumps(s, indent=2), encoding="utf-8")
            org.studio_session_id = s["session_id"]
            db.add(org)
        db.commit()
        if b.get("logo_png_base64"):
            save_logo(org.id, base64.b64decode(b["logo_png_base64"]))
        audit_append(event="operator_import", agent="human", team="operator",
                     human={"role": "operator", "action": "imported organization", "org_id": org.id},
                     output={"participants": len(people), "responses": len(b["responses"]), "reports": len(b["reports"])})
        return (f"Imported '{org.name}' as organization {org.id}. Sign in with username '{username}' and the password "
                f"you chose. {len(people)} participants (new random PINs), {len(b['responses'])} responses, "
                f"{len(b['entries'])} program values, {len(b['reports'])} report(s), {len(b['grants'])} grant draft(s).")


def report_data() -> str:
    """One compact JSON summary for the final report: per organization, tokens by task and what it built.
    Organization names only (no staff names, usernames, or participant data)."""
    import json

    from db import Indicator, IndicatorEntry, Participant, Response, StaffUser, TokenUsage

    with Session(engine()) as db:
        out = []
        for org in db.exec(select(Organization)).all():
            rows = db.exec(select(TokenUsage).where(TokenUsage.org_id == org.id)).all()
            tasks: dict = defaultdict(lambda: {"runs": 0, "requests": 0, "input": 0, "output": 0, "failed": 0})
            for r in rows:
                t = tasks[f"{r.feature}/{r.phase}"]
                t["runs"] += 1
                t["requests"] += r.requests
                t["input"] += r.input_tokens
                t["output"] += r.output_tokens
                t["failed"] += 0 if r.ok else 1
            count = lambda m: len(db.exec(select(m).where(m.org_id == org.id)).all())  # noqa: E731
            out.append({
                "org": org.name, "org_id": org.id, "created": org.created_at.strftime("%Y-%m-%d"),
                "staff": count(StaffUser), "studio_done": bool(org.measurement_plan),
                "survey_tools": count(Instrument), "participants": count(Participant), "responses": count(Response),
                "program_numbers": count(Indicator), "program_values": count(IndicatorEntry),
                "reports": count(ReportRun), "grants": count(GrantApplication),
                "tokens": {"input": sum(r.input_tokens for r in rows), "output": sum(r.output_tokens for r in rows),
                           "runs": len(rows)},
                "by_task": dict(sorted(tasks.items())),
            })
    return json.dumps({"through_line_report_data": out}, separators=(",", ":"))


def print_usage() -> None:
    with Session(engine()) as db:
        rows = ledger.platform_rows(db)
    print(f"{'Organization':34} {'Person':24} {'Requests':>9} {'Input':>11} {'Output':>10} {'Total':>11}")
    for org, person, req, inp, outp, tot in _sum(rows, lambda r: (r["organization"], r["person"])):
        print(f"{org[:34]:34} {person[:24]:24} {req:9,} {inp:11,} {outp:10,} {tot:11,}")
    grand = sum(r["total_tokens"] for r in rows)
    print(f"\n{len(rows)} agent runs, {grand:,} tokens in total.")


def print_orgs() -> None:
    with Session(engine()) as db:
        for org in db.exec(select(Organization)).all():
            parts = [f"{a['feature']}: {a['used']}/{a['limit']}" for a in allowances.summary(org) if a["used"] is not None]
            print(f"[{org.id}] {org.name}: " + "; ".join(parts))


def allow(org_id: int, kind: str, item: int | None, n: int) -> str:
    with Session(engine()) as db:
        org = db.get(Organization, org_id)
        if not org:
            raise SystemExit(f"No organization {org_id}.")
        if kind in ("studio", "report", "grants"):
            key = {"studio": "studio_starts", "report": "report_designs", "grants": "grants"}[kind]
            for _ in range(n):
                allowances.refund(org, key)
            db.add(org)
        else:
            model, field = {"report-revision": (ReportRun, "revisions"), "grant-revision": (GrantApplication, "revisions"),
                            "survey-review": (Instrument, "review_count")}[kind]
            row = db.get(model, item) if item else None
            if not row or row.org_id != org_id:
                raise SystemExit(f"No {kind.split('-')[0]} {item} for organization {org_id} (use --item).")
            setattr(row, field, max(0, (getattr(row, field) or 0) - n))
            db.add(row)
        db.commit()
        audit_append(event="operator_allowance", agent="human", team="operator",
                     human={"role": "operator", "action": f"allowed {n} more {kind}", "org_id": org_id},
                     output={"item": item})
        return f"Organization {org_id} ({org.name}) can use {n} more {kind}."


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("usage", help="token totals by organization and person, plus an Excel report")
    sub.add_parser("orgs", help="organizations and their allowance use")
    sub.add_parser("report-data", help="one-line JSON summary per organization (for the final report)")
    sub.add_parser("backfill", help="copy token use from the audit log for runs before the ledger existed")
    e = sub.add_parser("export-org", help="package an organization's work (no passwords or PINs) into a JSON file")
    e.add_argument("org_id", type=int)
    e.add_argument("out", type=Path)
    i = sub.add_parser("import-org", help="load a packaged organization under a new staff login")
    i.add_argument("bundle", type=Path)
    i.add_argument("--username", required=True)
    i.add_argument("--name", default="Demo administrator", help="display name for the new login")
    p = sub.add_parser("allow", help="give an organization more of a limited feature")
    p.add_argument("org_id", type=int)
    p.add_argument("kind", choices=["studio", "report", "grants", "report-revision", "grant-revision", "survey-review"])
    p.add_argument("--item", type=int, default=None)
    p.add_argument("--n", type=int, default=1)
    args = ap.parse_args()
    if args.cmd == "usage":
        print_usage()
        print(f"Excel report: {usage_report()}")
    elif args.cmd == "orgs":
        print_orgs()
    elif args.cmd == "report-data":
        print(report_data())
    elif args.cmd == "backfill":
        print(backfill())
    elif args.cmd == "export-org":
        print(export_org(args.org_id, args.out))
    elif args.cmd == "import-org":
        import getpass

        password = os.environ.get("IMPORT_PASSWORD") or getpass.getpass("Password for the new login (10+ characters, "
                                                                        "with a capital letter and a number): ")
        if len(password) < 10 or not any(ch.isupper() for ch in password) or not any(ch.isdigit() for ch in password):
            raise SystemExit("Please use at least 10 characters with a capital letter and a number.")
        print(import_org(args.bundle, args.username.strip().lower(), password, args.name))
    else:
        print(allow(args.org_id, args.kind, args.item, args.n))
