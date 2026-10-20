"""Play a fictional test profile through the real app, end to end, over HTTP.

    .venv/Scripts/python scripts/run_test_profile.py data/test_profiles/keystone_family_housing
    .venv/Scripts/python scripts/run_test_profile.py <profile> --from tools      # Studio already done
    .venv/Scripts/python scripts/run_test_profile.py <profile> --from report     # just rebuild the funder report
    .venv/Scripts/python scripts/run_test_profile.py <profile> --base http://127.0.0.1:8000

Stages: account -> studio -> tools -> collect -> numbers -> export -> report.
The Studio, Survey Review Bot, and Report Designer stages make LIVE gpt-6-luna calls when
the backend has a key. Only run them when Doug has asked for a live run.

Program numbers go in the way staff would: download each month's Excel template, fill in
the Value column, and upload it. The report stage saves the profile's brand + logo, picks
charts, and builds the funder deck and one-pager.

Each run writes data/test_profiles/<slug>/runs/<timestamp>/ (or --runs-dir) with a run log,
the participant codes/PINs created, the dashboard JSON, and every download.
Everything is synthetic; the same profile + seed gives the same participant data.
"""

import argparse
import io
import json
import random
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import httpx
from openpyxl import load_workbook

STAGES = ["account", "studio", "tools", "collect", "numbers", "export", "report"]
RATED_MAX = {"agree_5": 5, "confidence_5": 5, "frequency_5": 5, "scale_0_10": 10}
RATED_MIN = {"agree_5": 1, "confidence_5": 1, "frequency_5": 1, "scale_0_10": 0}
TIMEOUT = httpx.Timeout(900.0, connect=10.0)  # live agent phases can take minutes


def log(msg: str) -> None:
    print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)


def ok(r: httpx.Response) -> dict:
    if r.status_code >= 400:
        raise SystemExit(f"{r.request.method} {r.request.url} -> {r.status_code}: {r.text[:400]}")
    return r.json()


# ============================================================ interview answers


def pick_answer(question: str, bank: list[dict], used: set[str]) -> dict:
    q = question.lower()
    hits = lambda k: bool(re.search(r"(?<![a-z])" + re.escape(k) + r"(?![a-z])", q))  # whole words/phrases only
    scored = [(sum(hits(k) for k in item["keywords"]), -i, item) for i, item in enumerate(bank) if item["topic"] not in used]
    return max(scored, key=lambda t: (t[0], t[1]))[2] if scored else bank[-1]


# ============================================================ synthetic responses


def _spread(dist: dict[str, int], n: int, rng: random.Random) -> list[str]:
    values = [k for k, c in dist.items() for _ in range(c)]
    values = (values * (n // max(1, len(values)) + 1))[:n]
    rng.shuffle(values)
    return values


def make_people(spec: dict) -> list[dict]:
    rng = random.Random(spec["seed"])
    n = spec["count"]
    demo_cols = {k: _spread(v, n, rng) for k, v in spec["demographics"].items()}
    m = spec["response_model"]
    people = []
    for i in range(n):
        demo = {k: ([v[i]] if k == "race_ethnicity" else v[i]) for k, v in demo_cols.items()}
        people.append({"index": i, "improver": i < round(m["improver_share"] * n),
                       "skips_post": rng.random() < m["skip_post_share"], "demographics": demo})
    return people


def answer_value(q: dict, person: dict, wave: str, model: dict, reflections: list[str], rng: random.Random):
    scale = q["scale"]
    if scale in RATED_MAX:
        lo, hi = RATED_MIN[scale], RATED_MAX[scale]
        target = model["post_mean"] if (wave == "post" and person["improver"]) else model["pre_mean"]
        good = min(1.0, max(0.0, rng.gauss(target, model["noise"])))
        val = round(lo + good * (hi - lo))
        return hi + lo - val if q.get("reverse") else val
    if scale == "yes_no":
        p = model["yes_good_post"] if (wave == "post" and person["improver"]) else model["yes_good_pre"]
        good = 1 if rng.random() < p else 0
        return 1 - good if q.get("reverse") else good
    if scale == "number":
        return rng.randint(1, 12)
    if scale == "open_text":
        return reflections[person["index"] % len(reflections)] if wave == "post" else None
    return None


# ============================================================ the run


def fill_template(xlsx: bytes, values: dict[int, float]) -> bytes:
    """Type values into a downloaded program-numbers template, the way a staff member would."""
    wb = load_workbook(io.BytesIO(xlsx))
    ws = wb["Program numbers"]
    for row in ws.iter_rows(min_row=5):
        if row[0].value in values:
            row[5].value = values[row[0].value]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run(profile_dir: Path, base: str, start: str, runs_dir: Path | None = None) -> Path:
    profile = json.loads((profile_dir / "profile.json").read_text(encoding="utf-8"))
    out = (runs_dir or profile_dir / "runs") / datetime.now().strftime("%Y%m%d_%H%M%S")
    out.mkdir(parents=True)
    record: dict = {"profile": profile["slug"], "base": base, "started": datetime.now().isoformat(timespec="seconds"),
                    "stages": {}}
    staff = httpx.Client(base_url=base, timeout=TIMEOUT)
    todo = STAGES[STAGES.index(start):]
    if "account" not in todo:
        todo = ["account"] + todo  # always need to be signed in

    def save():
        (out / "run_log.json").write_text(json.dumps(record, indent=2, ensure_ascii=False), encoding="utf-8")

    # ---------------------------------------------------------------- account
    acct = profile["account"]
    r = staff.post("/api/auth/login", json={"username": acct["username"], "password": acct["password"]})
    if r.status_code == 401:
        r = staff.post("/api/auth/signup", json={**{k: acct[k] for k in ("username", "password", "display_name")},
                                                 "org_name": profile["org"]["name"]})
    me = ok(r)
    log(f"signed in as {acct['username']} ({me['org']['name']})")
    record["stages"]["account"] = {"org": me["org"]}

    # ---------------------------------------------------------------- studio (LIVE agents)
    if "studio" in todo:
        files = [("files", (Path(d).name, (profile_dir / d).read_bytes(), "application/pdf"))
                 for d in profile["org"]["documents"]]
        s = ok(staff.post("/api/logic-model/sessions", data={"org_name": profile["org"]["name"]}, files=files))
        base_url = f"/api/logic-model/sessions/{s['session_id']}"
        t0 = time.time()
        log("intake (team reads the documents, researches funders, proposes questions)…")
        s = ok(staff.post(f"{base_url}/intake"))
        used: set[str] = set()
        qa = []
        while s["phase"] == "interview":
            q = s["current_question"]["question"]
            item = pick_answer(q, profile["interview"]["answer_bank"], used)
            used.add(item["topic"])
            log(f"Q{s['question_number']}: {q[:90]}  -> answering '{item['topic']}'")
            qa.append({"question": q, "topic": item["topic"], "answer": item["answer"]})
            s = ok(staff.post(f"{base_url}/answer", json={"answer": item["answer"]}))
        log("drafting the logic model (draft, peer review, final)…")
        s = ok(staff.post(f"{base_url}/draft"))
        s = ok(staff.post(f"{base_url}/feedback", json={"feedback": profile["interview"]["logic_model_feedback"]}))
        s = ok(staff.post(f"{base_url}/approve"))
        log("measurement plan (research, draft, peer review, final)…")
        s = ok(staff.post(f"{base_url}/metrics/draft"))
        s = ok(staff.post(f"{base_url}/metrics/feedback", json={"feedback": profile["interview"]["metrics_feedback"]}))
        s = ok(staff.post(f"{base_url}/metrics/approve"))
        record["stages"]["studio"] = {"session_id": s["session_id"], "interview": qa, "logic_model": s["approved"],
                                      "plan": s["approved_plan"], "usage": s["usage"],
                                      "minutes": round((time.time() - t0) / 60, 1)}
        save()

    # ---------------------------------------------------------------- tools (+ LIVE Survey Review Bot)
    if "tools" in todo:
        tools = ok(staff.get("/api/tools"))
        if not tools["instruments"]:
            tools = ok(staff.post("/api/tools/build"))
        reviews = []
        for summary in tools["instruments"]:
            iid = summary["id"]
            log(f"Survey Review Bot: {summary['name']}…")
            d = ok(staff.post(f"/api/tools/instruments/{iid}/review"))
            applied = []
            qs = d["questions"]
            for issue in (d["review"] or {}).get("issues", []):
                fix = issue.get("fix")
                if not fix or issue["severity"] == "low" or fix["action"] == "hide":
                    continue
                if fix["action"] == "add_open_question" and fix.get("text"):
                    qs.append({"text": fix["text"], "scale": "open_text", "waves": ["post"], "source": "staff"})
                else:
                    q = next((x for x in qs if x["id"] == issue["question_id"]), None)
                    if not q:
                        continue
                    if fix["action"] == "rewrite" and fix.get("text"):
                        q["text"] = fix["text"]
                    elif fix["action"] == "change_scale" and fix.get("scale"):
                        q["scale"] = fix["scale"]
                    elif fix["action"] == "set_waves" and fix.get("waves"):
                        q["waves"] = fix["waves"]
                    elif fix["action"] == "set_reverse":
                        q["reverse"] = True
                applied.append({"action": fix["action"], "problem": issue["problem"]})
            demos = d["demographics"]
            if summary["method"] == "participant_survey":
                demos = list(profile["participants"]["demographics"].keys())
            d = ok(staff.put(f"/api/tools/instruments/{iid}", json={"name": d["name"], "questions": qs, "demographics": demos}))
            ok(staff.post(f"/api/tools/instruments/{iid}/publish"))
            reviews.append({"instrument": d["name"], "summary": (d["review"] or {}).get("summary"),
                            "ready": (d["review"] or {}).get("ready"), "applied": applied,
                            "remaining_checks": [c["problem"] for c in d["checks"]]})
        record["stages"]["tools"] = {"reviews": reviews}
        save()

    # ---------------------------------------------------------------- collect: participants + staff
    if "collect" in todo:
        spec = profile["participants"]
        people = make_people(spec)
        creds = ok(staff.post("/api/tools/participants", json={"count": spec["count"]}))["created"]
        (out / "participant_codes.json").write_text(json.dumps(creds, indent=2), encoding="utf-8")
        rng = random.Random(spec["seed"] + 1)
        tools = ok(staff.get("/api/tools"))["instruments"]
        counts = {"participant_pre": 0, "participant_post": 0, "staff_rows": 0}
        for person, cred in zip(people, creds):
            pc = httpx.Client(base_url=base, timeout=TIMEOUT)
            ok(pc.post("/api/auth/participant/login", json={"code": cred["code"], "pin": cred["pin"]}))
            for survey in ok(pc.get("/api/participant/surveys"))["surveys"]:
                for wave in ("pre", "post"):
                    if wave == "post" and person["skips_post"]:
                        continue
                    form = ok(pc.get(f"/api/participant/surveys/{survey['id']}/{wave}"))
                    detail = {q["id"]: q for q in ok(staff.get(f"/api/tools/instruments/{survey['id']}"))["questions"]}
                    answers = {q["id"]: v for q in form["questions"]
                               if (v := answer_value({**q, "reverse": detail[q["id"]]["reverse"]}, person, wave,
                                                     spec["response_model"], spec["reflections"], rng)) is not None}
                    demo = person["demographics"] if form["demographics"] else {}
                    ok(pc.post(f"/api/participant/surveys/{survey['id']}/{wave}", json={"answers": answers, "demographics": demo}))
                    counts[f"participant_{wave}"] += 1
            pc.close()
        for tool in [t for t in tools if t["method"] == "staff_assessment"]:
            detail = ok(staff.get(f"/api/tools/instruments/{tool['id']}"))
            for wave in ("pre", "post"):
                rows = []
                for person, cred in zip(people, creds):
                    if wave == "post" and person["skips_post"]:
                        continue
                    qs = [q for q in detail["questions"] if not q["hidden"] and wave in q["waves"]]
                    answers = {q["id"]: v for q in qs
                               if (v := answer_value(q, person, wave, spec["response_model"], spec["reflections"], rng)) is not None}
                    rows.append({"participant_code": cred["code"], "answers": answers})
                res = ok(staff.post(f"/api/tools/instruments/{tool['id']}/responses", json={"wave": wave, "rows": rows}))
                counts["staff_rows"] += res["saved"]
        log(f"collected: {counts}")
        record["stages"]["collect"] = counts
        save()

    # ---------------------------------------------------------------- program numbers
    if "numbers" in todo:
        pn = profile["program_numbers"]
        existing = {i["name"] for i in ok(staff.get("/api/tools/indicators"))["indicators"]}
        for c in pn["custom"]:
            if c["name"] not in existing:
                ok(staff.post("/api/tools/indicators", json={"name": c["name"], "unit": c["unit"], "description": ""}))
        inds = ok(staff.get("/api/tools/indicators"))["indicators"]
        plan_inds = [i for i in inds if i["source"] == "plan"]
        series, fallback = {}, iter(pn.get("plan_series", []))
        for ind in plan_inds:
            name = ind["name"].lower()
            rule = next((r for r in pn.get("plan_rules", []) if any(k in name for k in r["keywords"])), None)
            vals = rule["values"] if rule else next(fallback, None)
            if vals:
                series[ind["id"]] = vals
        for c in pn["custom"]:
            ind = next(i for i in inds if i["name"] == c["name"])
            series[ind["id"]] = c["values"]
        uploads = []
        for i, month in enumerate(pn["months"]):  # one template per month; uploading replaces that month
            r = staff.get("/api/tools/indicators/template.xlsx", params={"month": month})
            r.raise_for_status()
            filled = fill_template(r.content, {iid: vals[i] for iid, vals in series.items() if i < len(vals)})
            if i == len(pn["months"]) - 1:
                (out / f"program_numbers_{month}.xlsx").write_bytes(filled)  # keep one filled example
            res = ok(staff.post("/api/tools/indicators/upload",
                                files={"file": (f"program_numbers_{month}.xlsx", filled, "application/octet-stream")}))
            uploads.append({"month": month, "saved": res["saved"], "errors": res["errors"]})
        log(f"program numbers uploaded for {len(uploads)} months ({sum(u['saved'] for u in uploads)} values)")
        record["stages"]["numbers"] = {"indicators": len(series), "uploads": uploads}
        save()

    # ---------------------------------------------------------------- export
    if "export" in todo:
        dash = ok(staff.get("/api/dashboard"))
        (out / "dashboard.json").write_text(json.dumps(dash, indent=2, ensure_ascii=False), encoding="utf-8")
        for path, name in (("/api/dashboard/data.xlsx", "survey_data.xlsx"), ("/api/dashboard/report.pptx", "impact_report.pptx")):
            r = staff.get(path)
            r.raise_for_status()
            (out / name).write_bytes(r.content)
        record["stages"]["export"] = {"charts": [c["title"] for c in dash["charts"]], "headline": dash["headline"]}
        save()

    # ---------------------------------------------------------------- report (LIVE Report Designer)
    if "report" in todo:
        brand, req = profile["brand"], profile["report"]
        fields = {k: brand[k] for k in ("standards", "primary", "secondary", "accent", "heading_font", "body_font",
                                         "tagline", "contact")}
        logo = (profile_dir / brand["logo"]).read_bytes()
        opts = ok(staff.post("/api/reports/brand", data=fields, files={"logo": (brand["logo"], logo, "image/png")}))
        # Everything in the chosen groups except the survey-count chart, up to the API's limit of 20.
        chart_ids = [c["id"] for g in req["groups"] for c in opts["charts"] if c["group"] == g and c["id"] != "responses"][:20]
        log(f"Report Designer: {len(chart_ids)} charts, {req['minutes']} min for {req['audience']}…")
        t0 = time.time()
        built = ok(staff.post("/api/reports/build", json={"chart_ids": chart_ids, "audience": req["audience"],
                                                          "minutes": req["minutes"], "include_quotes": req["include_quotes"],
                                                          "notes": req["notes"]}))
        (out / "report_draft.json").write_text(json.dumps(built, indent=2, ensure_ascii=False), encoding="utf-8")
        for path, name in ((f"/api/reports/{built['id']}/deck.pptx", "funder_deck.pptx"),
                           (f"/api/reports/{built['id']}/one-pager.pdf", "impact_one_pager.pdf")):
            r = staff.get(path)
            r.raise_for_status()
            (out / name).write_bytes(r.content)
        p = built["privacy"]
        log(f"report built: {len(built['draft']['slides'])} slides; privacy clean={p['clean']} "
            f"(agent fixed {len(p['agent_findings'])}, code fixed {len(p['code_findings'])})")
        record["stages"]["report"] = {"report_id": built["id"], "charts": chart_ids, "slides": len(built["draft"]["slides"]),
                                      "privacy": p, "minutes": round((time.time() - t0) / 60, 1)}
    record["finished"] = datetime.now().isoformat(timespec="seconds")
    save()
    log(f"done -> {out}")
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("profile", type=Path)
    ap.add_argument("--base", default="http://127.0.0.1:8000")
    ap.add_argument("--from", dest="start", choices=STAGES, default="account")
    ap.add_argument("--runs-dir", type=Path, default=None, help="Where to write the run folder (default: <profile>/runs)")
    args = ap.parse_args()
    if not (args.profile / "profile.json").exists():
        sys.exit(f"No profile.json in {args.profile}")
    run(args.profile, args.base, args.start, args.runs_dir)
