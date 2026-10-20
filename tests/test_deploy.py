"""Live-site settings: the sign-up access code, serving the built website, and security headers."""

import sys
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).parent))
from test_prototype import STAFF  # noqa: E402


@pytest.fixture
def app(tmp_path, monkeypatch):
    monkeypatch.setenv("THROUGH_LINE_AUDIT", str(tmp_path / "audit_log.json"))
    monkeypatch.setenv("THROUGH_LINE_SESSIONS", str(tmp_path / "sessions"))
    monkeypatch.setenv("THROUGH_LINE_DB", str(tmp_path / "test.db"))
    import main
    import security

    security._fails.clear()
    return main.app


def test_new_organizations_need_the_access_code(app, monkeypatch):
    monkeypatch.setenv("THROUGH_LINE_SIGNUP_CODE", "harbor-lights-2026")
    c = TestClient(app)
    assert c.get("/api/health").json()["signup_code_required"] is True
    r = c.post("/api/auth/signup", json=STAFF)
    assert r.status_code == 403 and "access code" in r.json()["detail"]
    assert c.post("/api/auth/signup", json={**STAFF, "access_code": "wrong"}).status_code == 403
    assert c.post("/api/auth/signup", json={**STAFF, "access_code": "harbor-lights-2026"}).status_code == 200

    # Colleagues still join with an invite code, no access code needed.
    code = c.post("/api/auth/invite").json()["invite_code"]
    r = TestClient(app).post("/api/auth/signup", json={**STAFF, "username": "colleague", "invite_code": code})
    assert r.status_code == 200 and r.json()["user"]["role"] == "staff"


def test_open_signup_when_no_code_is_set(app, monkeypatch):
    monkeypatch.delenv("THROUGH_LINE_SIGNUP_CODE", raising=False)
    c = TestClient(app)
    assert c.get("/api/health").json()["signup_code_required"] is False
    assert c.post("/api/auth/signup", json=STAFF).status_code == 200


def test_serves_the_website_and_keeps_api_404s(tmp_path):
    import main

    dist = tmp_path / "dist"
    (dist / "assets").mkdir(parents=True)
    (dist / "index.html").write_text("<div id=root>Through Line</div>", encoding="utf-8")
    (dist / "assets" / "app.js").write_text("console.log(1)", encoding="utf-8")
    (dist / "favicon.svg").write_text("<svg/>", encoding="utf-8")
    (tmp_path / "secret.txt").write_text("nope", encoding="utf-8")
    site = FastAPI()
    main.mount_frontend(site, dist)
    c = TestClient(site)
    assert "Through Line" in c.get("/").text
    assert "Through Line" in c.get("/grants").text  # client-side routes get the app
    assert c.get("/assets/app.js").text == "console.log(1)"
    assert c.get("/favicon.svg").text == "<svg/>"
    assert "nope" not in c.get("/../secret.txt").text and "nope" not in c.get("/%2e%2e/secret.txt").text
    assert c.get("/api/nothing-here").status_code == 404


def test_security_headers(app):
    r = TestClient(app).get("/api/health")
    assert r.headers["x-content-type-options"] == "nosniff" and r.headers["x-frame-options"] == "DENY"
    assert r.headers["cache-control"] == "no-store"


# ============================================================ protections from the pre-launch review

def test_oversized_requests_are_refused_before_reading(app, monkeypatch):
    import main

    monkeypatch.setattr(main, "MAX_REQUEST_BYTES", 1000)
    r = TestClient(app).post("/api/auth/signup", content=b"x" * 2000, headers={"content-type": "application/json"})
    assert r.status_code == 413


def test_sign_in_attempts_are_limited_per_address(app):
    c = TestClient(app)
    codes = [c.post("/api/auth/login", json={"username": f"nobody{i}", "password": "Wrong-Pass-1234"}).status_code
             for i in range(32)]
    assert codes[0] == 401 and codes[-1] == 429


def test_audit_log_never_keeps_flagged_text():
    from agents.common import _dump
    from models import PrivacyReview

    review = PrivacyReview(clean=False, summary="One name found.", findings=[
        {"where": "slide 3", "text": "Maria from Elm Street", "issue": "a name", "replacement": "a parent"}])
    logged = str(_dump(review))
    assert "Maria" not in logged and "[withheld]" in logged and "slide 3" in logged
    assert "203-555-0182" not in str(_dump({"note": "call 203-555-0182"}))


def test_bad_numbers_and_sheet_names_cant_break_an_org(org_env):
    from db import Instrument, engine
    from sqlmodel import Session, select

    c, _ = org_env
    ind = c.post("/api/tools/indicators", json={"name": "=HYPERLINK(\"http://x\")", "unit": "$"}).json()["indicators"][0]
    assert c.post(f"/api/tools/indicators/{ind['id']}/entries",
                  content=b'{"entries": [{"period": "2026-09-01", "value": NaN, "note": ""}]}',
                  headers={"content-type": "application/json"}).status_code == 422
    with Session(engine()) as db:
        inst = db.exec(select(Instrument)).first()
        inst.name = "How are you? [2026]"
        db.add(inst)
        db.commit()
    assert c.get("/api/dashboard/data.xlsx").status_code == 200
    from openpyxl import load_workbook
    import io as _io

    wb = load_workbook(_io.BytesIO(c.get("/api/tools/indicators/template.xlsx", params={"month": "2026-09"}).content))
    names = [row[1] for row in wb["Program numbers"].iter_rows(min_row=5, values_only=True) if row[1]]
    assert all(not str(n).startswith("=") for n in names)  # staff text can't become a live formula


def test_private_addresses_are_never_fetched():
    import intake

    for url in ("http://127.0.0.1:8000", "http://localhost", "http://169.254.169.254/latest/meta-data"):
        try:
            intake.read_website(url)
            raise AssertionError(f"fetched {url}")
        except intake.IntakeError as exc:
            assert "private network" in str(exc)
