"""Hardening for the public site: the shared demo account is view-only and can't be locked, and hostile
Office files (zip bombs, sheets that claim a million rows) are refused or read only as far as a real template goes."""

import io
import time
import zipfile

import pytest


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("THROUGH_LINE_AUDIT", str(tmp_path / "audit_log.json"))
    monkeypatch.setenv("THROUGH_LINE_SESSIONS", str(tmp_path / "sessions"))
    monkeypatch.setenv("THROUGH_LINE_DB", str(tmp_path / "test.db"))
    from fastapi.testclient import TestClient

    import main

    return TestClient(main.app)


DEMO = {"username": "keystone.demo", "password": "Shared-Demo-2026", "display_name": "Demo", "org_name": "Demo Org"}


def test_demo_account_is_view_only(client):
    assert client.post("/api/auth/signup", json=DEMO).status_code == 200
    assert client.get("/api/auth/me").json()["demo"] is True
    assert client.get("/api/dashboard").status_code == 200  # looking and downloading still work
    r = client.post("/api/tools/build")
    assert r.status_code == 403 and "view-only" in r.json()["detail"]
    assert client.post("/api/grants", json={}).status_code == 403
    assert client.post("/api/auth/logout").status_code == 200


def test_ordinary_accounts_are_not_view_only(client):
    staff = {**DEMO, "username": "someone.else", "org_name": "Other Org"}
    assert client.post("/api/auth/signup", json=staff).status_code == 200
    assert client.get("/api/auth/me").json()["demo"] is False
    assert client.post("/api/tools/build").status_code != 403


def test_wrong_passwords_never_lock_the_shared_demo(client):
    assert client.post("/api/auth/signup", json=DEMO).status_code == 200
    client.post("/api/auth/logout")
    for _ in range(7):
        assert client.post("/api/auth/login", json={"username": "keystone.demo", "password": "wrong-guess"}).status_code == 401
    ok = client.post("/api/auth/login", json={"username": "keystone.demo", "password": DEMO["password"]})
    assert ok.status_code == 200


def _zip_bomb() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", b"\0" * (150 * 1024 * 1024))  # ~150 MB that compresses to almost nothing
    return buf.getvalue()


def test_zip_bombs_are_refused_before_unpacking():
    import intake
    import program_numbers

    bomb = _zip_bomb()
    assert len(bomb) < 1024 * 1024
    with pytest.raises(intake.IntakeError):
        intake.read_upload(1, "plan.docx", bomb)
    with pytest.raises(ValueError):
        program_numbers.parse_upload(bomb)


def test_a_sheet_claiming_a_million_rows_is_read_quickly():
    from openpyxl import Workbook

    import program_numbers

    wb = Workbook()
    ws = wb.active
    ws.title = "Program numbers"
    ws["A1"] = "Number id"
    ws["XFD1048576"] = "far away"  # declares a huge sheet with almost nothing in it
    buf = io.BytesIO()
    wb.save(buf)
    start = time.time()
    rows, _errors = program_numbers.parse_upload(buf.getvalue())
    assert rows == [] and time.time() - start < 10
