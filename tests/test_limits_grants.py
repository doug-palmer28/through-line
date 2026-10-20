"""Offline tests for Step 15b: usage limits, the token ledger, live numbers in the report shell, and grants."""

import io
import json
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pydantic_ai.models.test import TestModel

sys.path.insert(0, str(Path(__file__).parent))
from test_prototype import STAFF, _with_plan  # noqa: E402
from test_visible_text import _pdf  # noqa: E402  (a PDF with "SECRET" hidden six ways)


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setenv("THROUGH_LINE_AUDIT", str(tmp_path / "audit_log.json"))
    monkeypatch.setenv("THROUGH_LINE_SESSIONS", str(tmp_path / "sessions"))
    monkeypatch.setenv("THROUGH_LINE_DB", str(tmp_path / "test.db"))
    monkeypatch.setenv("THROUGH_LINE_BRAND_DIR", str(tmp_path / "brand"))
    from agents import common
    import main
    import security

    security._fails.clear()
    monkeypatch.setattr(common, "MODEL_OVERRIDE", TestModel())
    c = TestClient(main.app)
    assert c.post("/api/auth/signup", json=STAFF).status_code == 200
    _with_plan(main.app, c)
    c.post("/api/tools/build")
    return c, tmp_path


def _log(tmp) -> list[dict]:
    return json.loads((tmp / "audit_log.json").read_text(encoding="utf-8"))


def _charts(c) -> list[str]:
    return [ch["id"] for ch in c.get("/api/reports").json()["charts"]][:3]


def _set_draft(model, row_id: int, mutate) -> None:
    from db import engine
    from sqlmodel import Session

    with Session(engine()) as db:
        row = db.get(model, row_id)
        draft = json.loads(json.dumps(row.draft))
        mutate(draft)
        row.draft = draft
        db.add(row)
        db.commit()


# ============================================================ allowances


def test_studio_starts_are_limited(env):
    c, _ = env
    for _ in range(2):
        assert c.post("/api/logic-model/sessions", data={"org_name": "Harbor Light Housing"}).status_code == 200
    r = c.post("/api/logic-model/sessions", data={"org_name": "Harbor Light Housing"})
    assert r.status_code == 429 and "Studio sessions" in r.json()["detail"]


def test_survey_review_bot_is_limited_per_tool(env):
    c, _ = env
    iid = c.get("/api/tools").json()["instruments"][0]["id"]
    for left in (1, 0):
        r = c.post(f"/api/tools/instruments/{iid}/review")
        assert r.status_code == 200 and r.json()["reviews_left"] == left
    assert c.post(f"/api/tools/instruments/{iid}/review").status_code == 429


def test_report_shell_designs_once_revises_once_and_refreshes_free(env):
    from db import ReportRun

    c, tmp = env
    run = c.post("/api/reports/build", json={"chart_ids": _charts(c)}).json()
    assert c.get("/api/reports").json()["designs_left"] == 0
    assert c.post("/api/reports/build", json={"chart_ids": _charts(c)}).status_code == 429

    # Numbers are tokens: refreshing fills in the latest value with no AI call.
    def tokens(d):
        d["deck_title"] = "{{participants}} families enrolled"
        d["one_pager"]["stats"] = [{"value": "{{participants}}", "label": "families"},
                                   {"value": "{{not_a_number}}", "label": "x"}]
    _set_draft(ReportRun, run["id"], tokens)
    before = c.post(f"/api/reports/{run['id']}/refresh").json()
    assert before["draft"]["deck_title"] == "0 families enrolled"
    assert before["missing_numbers"] == ["not_a_number"] and before["draft"]["one_pager"]["stats"][1]["value"] == "–"
    c.post("/api/tools/participants", json={"count": 7})
    after = c.post(f"/api/reports/{run['id']}/refresh").json()
    assert after["draft"]["deck_title"] == "7 families enrolled"
    agents_before = sum(e["event"] == "agent_finished" for e in _log(tmp))
    deck = c.get(f"/api/reports/{run['id']}/deck.pptx")
    assert deck.status_code == 200 and b"7 families enrolled" in _pptx_text(deck.content)
    assert sum(e["event"] == "agent_finished" for e in _log(tmp)) == agents_before  # downloads never call the AI

    r = c.post(f"/api/reports/{run['id']}/revise", json={"feedback": "Lead with housing stability."})
    assert r.status_code == 200 and r.json()["revisions_left"] == 0 and r.json()["feedback"][0]["text"]
    assert c.post(f"/api/reports/{run['id']}/revise", json={"feedback": "Again"}).status_code == 429

    # The privacy check: re-checking an unchanged report is free (no agent run); a change triggers a real check.
    agents_before = sum(e["event"] == "agent_finished" for e in _log(tmp))
    assert c.post(f"/api/reports/{run['id']}/privacy-check").status_code == 200
    assert sum(e["event"] == "agent_finished" for e in _log(tmp)) == agents_before
    _set_draft(ReportRun, run["id"], lambda d: d.update(deck_title="A changed title"))  # it now reads differently
    for _ in range(2):
        assert c.post(f"/api/reports/{run['id']}/privacy-check").status_code == 200
    assert sum(e["event"] == "agent_finished" for e in _log(tmp)) == agents_before + 1  # checked once, then cached
    events = {e["event"] for e in _log(tmp)}
    assert {"report_built", "report_revised", "report_refreshed", "privacy_checked"} <= events


def _pptx_text(data: bytes) -> bytes:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    return " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame).encode()


def test_failed_ai_run_gives_the_allowance_back(env, monkeypatch):
    import routes.reports as rr

    c, _ = env

    async def boom(*a, **k):
        raise RuntimeError("model unavailable")

    monkeypatch.setattr(rr, "run_agent", boom)
    assert c.post("/api/reports/build", json={"chart_ids": _charts(c)}).status_code == 502
    assert c.get("/api/reports").json()["designs_left"] == 1


# ============================================================ token ledger


def test_every_run_is_recorded_per_person(env):
    from db import TokenUsage, engine
    from sqlmodel import Session, select

    c, _ = env
    c.post("/api/reports/build", json={"chart_ids": _charts(c)})
    with Session(engine()) as db:
        rows = db.exec(select(TokenUsage)).all()
    assert {r.agent for r in rows} == {"report_designer", "privacy_reviewer"}
    assert all(r.feature == "impact_report" and r.org_id and r.user_id and r.input_tokens > 0 for r in rows)

    u = c.get("/api/usage").json()
    assert u["me"]["total_tokens"] == u["organization"]["total_tokens"] > 0
    assert u["organization"]["by_person"][0]["key"] == STAFF["display_name"]  # admins see who used what
    assert any(a["feature"].startswith("Impact report") and a["used"] == 1 for a in u["allowances"])

    import main
    other = TestClient(main.app)
    other.post("/api/auth/signup", json={**STAFF, "username": "another_org", "org_name": "Other Org"})
    assert other.get("/api/usage").json()["organization"]["total_tokens"] == 0  # orgs never see each other


def test_operator_report_lists_every_org(env, tmp_path):
    c, _ = env
    c.post("/api/reports/build", json={"chart_ids": _charts(c)})
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
    import operator_tools

    out = operator_tools.usage_report(tmp_path / "reports")
    from openpyxl import load_workbook

    wb = load_workbook(out)
    assert {"By organization", "By person", "By month", "Every run"} <= set(wb.sheetnames)
    rows = list(wb["By person"].iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == STAFF["org_name"] and rows[0][1] == STAFF["display_name"] and rows[0][-1] > 0


# ============================================================ grants


def _spy(monkeypatch) -> list[str]:
    """Capture every prompt any agent receives."""
    import agents.funder_researcher as fr
    import agents.privacy_reviewer as pr
    import routes.grants as rg

    seen: list[str] = []
    for module in (fr, pr, rg):
        real = module.run_agent

        async def spy(agent, name, prompt, *a, _real=real, **k):
            seen.append(prompt)
            return await _real(agent, name, prompt, *a, **k)

        monkeypatch.setattr(module, "run_agent", spy)
    return seen


def _apply(c, foundation="Elm City Family Fund", files=None):
    return c.post("/api/grants", data={"foundation": foundation, "opportunity": "Housing Stability RFP 2027",
                                       "questions": "1. Describe your program (250 words)\n2. How do you measure results?"},
                  files=files or [])


def test_grant_application_end_to_end(env, monkeypatch):
    from db import GrantApplication

    c, tmp = env
    seen = _spy(monkeypatch)
    r = _apply(c, files=[("files", ("rfp.pdf", _pdf(), "application/pdf"))])
    assert r.status_code == 200, r.text
    app = r.json()
    assert app["draft"]["answers"] and app["profile"]["sources"] is not None
    assert app["research_mode"] == "knowledge"  # TestModel can't web search: the fallback marks sources unverified
    assert all(not s["verified"] and s["url"] is None for s in app["profile"]["sources"])
    assert app["sources"][0]["hidden"]  # the RFP's hidden text was counted...
    assert seen and all("SECRET" not in p for p in seen)  # ...and never reached any agent
    assert "Our mission is stable homes" in seen[0]  # while the visible text did

    finished = [e["agent"] for e in _log(tmp) if e["event"] == "agent_finished"]
    assert {"funder_researcher", "grant_writer", "privacy_reviewer"} <= set(finished)
    assert {"grant_built", "hidden_text_skipped"} <= {e["event"] for e in _log(tmp)}

    # Live numbers in the application, filled in the Word file
    _set_draft(GrantApplication, app["id"], lambda d: d["answers"][0].update(
        answer="We have enrolled {{participants}} families. [ADD: total program budget]", word_limit=5))
    view = c.get(f"/api/grants/{app['id']}").json()
    assert view["draft"]["answers"][0]["answer"].startswith("We have enrolled 0 families")
    assert view["draft"]["answers"][0]["over_limit"] is True
    doc = c.get(f"/api/grants/{app['id']}/application.docx")
    assert doc.status_code == 200 and doc.content[:2] == b"PK"
    from docx import Document
    text = "\n".join(p.text for p in Document(io.BytesIO(doc.content)).paragraphs)
    assert "We have enrolled 0 families" in text and "[ADD: total program budget]" in text

    assert c.post(f"/api/grants/{app['id']}/revise", json={"feedback": "Shorter."}).json()["revisions_left"] == 0
    assert c.post(f"/api/grants/{app['id']}/revise", json={"feedback": "Again."}).status_code == 429
    assert c.post(f"/api/grants/{app['id']}/privacy-check").status_code == 200


def test_grants_are_limited_per_month_and_private(env):
    import main

    c, _ = env
    ids = [_apply(c, f"Fund {i}").json()["id"] for i in range(3)]
    assert c.get("/api/grants").json()["left_this_month"] == 0
    r = _apply(c, "Fund 4")
    assert r.status_code == 429 and "this month" in r.json()["detail"]

    other = TestClient(main.app)
    other.post("/api/auth/signup", json={**STAFF, "username": "someone_else", "org_name": "Other Org"})
    assert other.get(f"/api/grants/{ids[0]}").status_code == 404
    assert other.get(f"/api/grants/{ids[0]}/application.docx").status_code == 404
    assert other.get("/api/grants").json()["applications"] == []
    assert TestClient(main.app).get("/api/grants").status_code == 401


def test_grant_rejects_bad_uploads_without_spending(env):
    c, _ = env
    r = _apply(c, files=[("files", ("x.exe", b"MZ", "application/octet-stream"))])
    assert r.status_code == 400
    assert c.get("/api/grants").json()["left_this_month"] == 3
