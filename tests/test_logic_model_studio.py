"""Offline end-to-end tests of the Studio (logic model + measurement plan).

Uses pydantic-ai's TestModel instead of gpt-6-luna, so no live model calls are
made. Audit log and session records go to a temp folder, never the real ones.
"""

import asyncio
import io
import json
from datetime import timedelta
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation
from pydantic_ai.models.test import TestModel

MARKER = "ZEBRA-PRIVATE-MARKER-7731"  # appears only inside uploaded files


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("THROUGH_LINE_AUDIT", str(tmp_path / "audit_log.json"))
    monkeypatch.setenv("THROUGH_LINE_SESSIONS", str(tmp_path / "sessions"))
    monkeypatch.setenv("THROUGH_LINE_DB", str(tmp_path / "test.db"))
    from agents import common
    import main
    import sessions

    sessions._SESSIONS.clear()
    monkeypatch.setattr(common, "MODEL_OVERRIDE", TestModel())
    c = TestClient(main.app)
    r = c.post("/api/auth/signup", json={"username": "ed_harbor", "password": "Harbor-Light-2026",
                                         "display_name": "Test ED", "org_name": "Harbor Light Housing"})
    assert r.status_code == 200, r.text
    return c


def _docx_bytes(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _start(client, **files) -> str:
    r = client.post(
        "/api/logic-model/sessions",
        data={"org_name": "Harbor Light Housing"},
        files=[("files", ("deck.docx", _docx_bytes(f"Harbor Light helps families. {MARKER}"),
                          "application/octet-stream")),
               ("files", ("notes.txt", b"We run a 12-week case management program.", "text/plain"))],
    )
    assert r.status_code == 200, r.text
    assert r.json()["phase"] == "intake"
    sid = r.json()["session_id"]
    r = client.post(f"/api/logic-model/sessions/{sid}/intake")
    assert r.status_code == 200, r.text
    return sid


def _record(tmp_path, sid) -> str:
    return (tmp_path / "sessions" / f"{sid}.json").read_text(encoding="utf-8")


def test_full_flow(client, tmp_path):
    base = "/api/logic-model/sessions"
    sid = _start(client)
    view = client.get(f"{base}/{sid}").json()
    assert view["phase"] == "interview" and view["question_number"] == 1
    assert [s["kind"] for s in view["sources"]] == ["docx", "txt"]
    assert MARKER not in json.dumps(view)

    import sessions
    assert sessions.get(sid).source_text is None and sessions.get(sid).pending_sources == []
    assert sessions.get(sid).sector_brief is not None

    # Live feed shows the team working (handoff to the sector expert, its reply, done)
    events = client.get(f"{base}/{sid}/feed").json()["events"]
    kinds = {e["kind"] for e in events}
    assert {"thinking", "ask", "reply", "done"} <= kinds
    assert client.get(f"{base}/{sid}/feed", params={"after": events[-1]["seq"]}).json()["events"] == []

    # Five answers -> ready to draft
    for n in range(1, 6):
        r = client.post(f"{base}/{sid}/answer", json={"answer": f"Answer number {n}"})
        assert r.status_code == 200, r.text
    assert r.json()["phase"] == "ready_to_draft"

    # Draft, direct edit, one revision, approve
    view = client.post(f"{base}/{sid}/draft").json()
    assert view["phase"] == "review" and view["draft"]["version"] == 1
    model = view["draft"]["logic_model"]
    model["activities"] = model["activities"] + ["Weekly peer mentoring circles"]
    view = client.put(f"{base}/{sid}/logic-model", json=model).json()
    assert view["draft"]["source"] == "client_edit"
    assert "Weekly peer mentoring circles" in view["draft"]["logic_model"]["activities"]
    view = client.post(f"{base}/{sid}/feedback", json={"feedback": "Add peer mentoring."}).json()
    assert view["draft"]["version"] == 3 and view["revisions_left"] == 1
    view = client.post(f"{base}/{sid}/approve").json()
    assert view["phase"] == "ready_for_metrics"

    # Measurement plan: draft, edit, revise, approve
    view = client.post(f"{base}/{sid}/metrics/draft").json()
    assert view["phase"] == "metrics_review" and view["metric_draft"]["version"] == 1
    plan = view["metric_draft"]["plan"]
    plan["metrics"][0]["name"] = "Housing stability confidence"
    view = client.put(f"{base}/{sid}/metrics", json=plan).json()
    assert view["metric_draft"]["source"] == "client_edit"
    view = client.post(f"{base}/{sid}/metrics/feedback", json={"feedback": "Shorter survey please."}).json()
    assert view["metric_draft"]["version"] == 3 and view["metric_revisions_left"] == 1
    view = client.post(f"{base}/{sid}/metrics/approve").json()
    assert view["phase"] == "complete" and view["approved_plan"]

    # Download includes the plan slides
    r = client.get(f"{base}/{sid}/export.pptx")
    assert r.status_code == 200
    titles = [sh.text_frame.text for sl in Presentation(io.BytesIO(r.content)).slides for sh in sl.shapes
              if sh.has_text_frame]
    assert any("MEASUREMENT PLAN" in t for t in titles)

    # Saved record: typed answers yes, uploads no, working notes wiped at completion
    record = _record(tmp_path, sid)
    assert "Answer number 3" in record and "Shorter survey please." in record
    assert MARKER not in record
    assert "working_notes" not in json.loads(record)
    assert sessions.get(sid).snapshot is None

    # Audit: chain intact, key events present, no uploaded content
    audit = client.get(f"{base}/{sid}/audit").json()
    assert audit["verify"]["ok"], audit["verify"]
    events = {e["event"] for e in audit["entries"]}
    for expected in ["session_started", "run_started", "agent_started", "agent_finished", "handoff",
                     "sources_discarded", "interview_complete", "human_edited", "human_approved",
                     "working_notes_wiped", "export_created", "run_finished"]:
        assert expected in events, expected
    assert {e["model"] for e in audit["entries"] if e["agent"] == "boss"} == {"gpt-6-luna"}
    assert MARKER not in (tmp_path / "audit_log.json").read_text(encoding="utf-8")

    assert client.delete(f"{base}/{sid}").json() == {"ended": True}
    assert client.get(f"{base}/{sid}").status_code == 404


def test_session_survives_server_restart(client, tmp_path):
    import sessions

    sid = _start(client)
    client.post(f"/api/logic-model/sessions/{sid}/answer", json={"answer": "First answer"})
    sessions._SESSIONS.clear()  # what a server restart does to memory
    view = client.get(f"/api/logic-model/sessions/{sid}").json()
    assert view["answered"] == 1 and view["phase"] == "interview"
    assert sessions.get(sid).snapshot is not None  # working notes restored for the active session
    r = client.post(f"/api/logic-model/sessions/{sid}/answer", json={"answer": "Second answer"})
    assert r.status_code == 200 and r.json()["answered"] == 2


def test_idle_session_expires_and_wipes_notes(client, tmp_path):
    import sessions

    sid = _start(client)
    sessions.get(sid).last_active -= sessions.SESSION_LIFETIME + timedelta(minutes=1)
    assert client.get(f"/api/logic-model/sessions/{sid}").status_code == 404
    record = json.loads(_record(tmp_path, sid))
    assert record["ended"] and "working_notes" not in record


def test_word_limit_enforced(client):
    sid = _start(client)
    r = client.post(f"/api/logic-model/sessions/{sid}/answer", json={"answer": "word " * 201})
    assert r.status_code == 422


def test_extra_document_is_read_not_saved(client, tmp_path):
    sid = _start(client)
    r = client.post(f"/api/logic-model/sessions/{sid}/documents",
                    files={"file": ("report.docx", _docx_bytes(f"Annual report {MARKER}"), "application/octet-stream")})
    assert r.status_code == 200, r.text
    view = r.json()
    assert view["uploads_left"] == 1 and view["upload_request"] is None
    assert len(view["sources"]) == 3
    assert MARKER not in _record(tmp_path, sid)
    assert MARKER not in (tmp_path / "audit_log.json").read_text(encoding="utf-8")
    assert list((tmp_path / "sessions").iterdir()) == [tmp_path / "sessions" / f"{sid}.json"]


def test_consult_refused_when_budget_spent(client, tmp_path):
    import sessions
    from agents.boss import consult
    from agents.common import TeamDeps

    deps = TeamDeps(session=sessions.create("Refusal Org"), run_id="run-test", phase="interview", consults_left=0)
    reply = asyncio.run(consult(SimpleNamespace(deps=deps), "sector_expert", "anything?"))
    assert reply.startswith("Refused")
    assert deps.consult_log == []
    log = json.loads((tmp_path / "audit_log.json").read_text(encoding="utf-8"))
    assert log[-1]["event"] == "limit_reached"


def test_rejects_bad_files_and_private_urls(client):
    r = client.post("/api/logic-model/sessions", data={"org_name": "X"},
                    files=[("files", ("virus.exe", b"MZ", "application/octet-stream"))])
    assert r.status_code == 400
    r = client.post("/api/logic-model/sessions", data={"org_name": "X", "website_url": "http://127.0.0.1:8000"})
    assert r.status_code == 400


def test_phase_order_enforced(client):
    sid = _start(client)
    assert client.post(f"/api/logic-model/sessions/{sid}/metrics/draft").status_code == 409
    assert client.post(f"/api/logic-model/sessions/{sid}/approve").status_code == 409
    assert client.post("/api/logic-model/sessions/ses-nope/answer", json={"answer": "hi"}).status_code == 404
