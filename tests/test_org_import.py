"""The Keystone demo bundle loads into a fresh database (as on the live site) and everything works offline."""

import io
import sys
from pathlib import Path

from fastapi.testclient import TestClient
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "scripts" / "demo" / "keystone_family_housing.json"


def test_keystone_bundle_imports_and_works(tmp_path, monkeypatch):
    for var, name in (("THROUGH_LINE_AUDIT", "audit_log.json"), ("THROUGH_LINE_SESSIONS", "sessions"),
                      ("THROUGH_LINE_DB", "fresh.db"), ("THROUGH_LINE_BRAND_DIR", "brand")):
        monkeypatch.setenv(var, str(tmp_path / name))
    sys.path.insert(0, str(ROOT / "scripts"))
    import main
    import operator_tools

    # Someone already signed up on the live site first: the import must not collide with their ids.
    TestClient(main.app).post("/api/auth/signup", json={"username": "early_bird", "password": "Early-Bird-2026",
                                                         "display_name": "E", "org_name": "First Org"})
    msg = operator_tools.import_org(BUNDLE, "keystone.demo", "Keystone-Demo-2026", "Demo ED")
    assert "Imported 'Keystone Family Housing Partners'" in msg

    c = TestClient(main.app)
    me = c.post("/api/auth/login", json={"username": "keystone.demo", "password": "Keystone-Demo-2026"}).json()
    assert me["org"]["name"] == "Keystone Family Housing Partners" and me["user"]["role"] == "admin"
    assert c.get(f"/api/logic-model/sessions/{me['studio_session_id']}").json()["phase"] == "complete"

    dash = c.get("/api/dashboard").json()
    assert dash["headline"]["participants"] == 24 and dash["headline"]["matched"] == 20
    assert any(ch["id"].startswith("indicator-") for ch in dash["charts"])

    report = c.get("/api/reports").json()
    assert report["has_logo"] and report["runs"] and report["designs_left"] == 0
    run = report["runs"][0]
    deck = c.get(f"/api/reports/{run['id']}/deck.pptx")
    slides = Presentation(io.BytesIO(deck.content)).slides
    chart_slides = sum(any(sh.has_chart for sh in s.shapes) for s in slides)
    assert deck.status_code == 200 and chart_slides >= 3  # chart ids were re-pointed at the new rows
    assert c.get(f"/api/reports/{run['id']}/one-pager.pdf").content.startswith(b"%PDF")
    assert c.get("/api/grants").json()["applications"]

    # Nobody knows the participants' PINs, including from the old run's participant_codes.json.
    code = dash and c.get("/api/tools/participants").json()["participants"][0]["code"]
    assert TestClient(main.app).post("/api/auth/participant/login", json={"code": code, "pin": "000000"}).status_code == 401
