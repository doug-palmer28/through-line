"""Offline tests for the working prototype: accounts, portals, survey tools, program numbers, dashboard, downloads."""

import io
import json
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation
from pydantic_ai.models.test import TestModel

sys.path.insert(0, str(Path(__file__).parent))
import demo_server  # noqa: E402  (provides a realistic approved plan)

STAFF = {"username": "ed_harbor", "password": "Harbor-Light-2026", "display_name": "Test ED",
         "org_name": "Harbor Light Housing"}


@pytest.fixture
def app_env(tmp_path, monkeypatch):
    monkeypatch.setenv("THROUGH_LINE_AUDIT", str(tmp_path / "audit_log.json"))
    monkeypatch.setenv("THROUGH_LINE_SESSIONS", str(tmp_path / "sessions"))
    monkeypatch.setenv("THROUGH_LINE_DB", str(tmp_path / "test.db"))
    from agents import common
    import main
    import security

    security._fails.clear()
    monkeypatch.setattr(common, "MODEL_OVERRIDE", TestModel())
    return main.app, tmp_path


def _staff(app, **overrides) -> TestClient:
    c = TestClient(app)
    r = c.post("/api/auth/signup", json={**STAFF, **overrides})
    assert r.status_code == 200, r.text
    return c


def _with_plan(app, client) -> None:
    """Give the org an approved logic model + plan (as the Studio would on approval)."""
    from db import Organization, engine
    from sqlmodel import Session

    org_id = client.get("/api/auth/me").json()["org"]["id"]
    plan = json.loads(json.dumps(demo_server.CANNED["MetricsResult"]["plan"]))
    plan["metrics"].append({"name": "Families served", "outcome": "Families enrolled", "method": "program_records",
                            "questions": [{"text": "Families enrolled this month", "scale": "number", "included": True,
                                           "source_ids": []}],
                            "timing": "Monthly", "how_scored": "Count", "why_it_matters": "Reach",
                            "funder_relevance": "", "source_ids": []})
    with Session(engine()) as db:
        org = db.get(Organization, org_id)
        org.logic_model = demo_server.LOGIC_MODEL
        org.measurement_plan = plan
        db.add(org)
        db.commit()


# ============================================================ accounts and security


def test_accounts_are_hashed_and_protected(app_env):
    app, tmp = app_env
    c = _staff(app)
    assert c.get("/api/auth/me").json()["org"]["name"] == "Harbor Light Housing"

    # Nothing sensitive is stored in plain text
    raw = (tmp / "test.db").read_bytes()
    assert b"ed_harbor" not in raw and b"Harbor-Light-2026" not in raw
    from db import StaffUser, engine
    from sqlmodel import Session, select
    with Session(engine()) as db:
        user = db.exec(select(StaffUser)).one()
        assert user.password_hash.startswith("$argon2id$") and len(user.username_hash) == 64

    # Cookie is HttpOnly + SameSite=Strict
    r = TestClient(app).post("/api/auth/login", json={"username": "ED_Harbor", "password": STAFF["password"]})
    cookie = r.headers["set-cookie"].lower()
    assert r.status_code == 200 and "httponly" in cookie and "samesite=strict" in cookie

    # Logged out: everything private is closed
    anon = TestClient(app)
    for path in ["/api/auth/me", "/api/tools", "/api/dashboard", "/api/dashboard/data.xlsx"]:
        assert anon.get(path).status_code == 401, path
    assert anon.post("/api/logic-model/sessions", data={"org_name": "X"}).status_code == 401

    # Weak password, duplicate username
    assert TestClient(app).post("/api/auth/signup", json={**STAFF, "username": "other", "password": "short"}).status_code == 422
    assert TestClient(app).post("/api/auth/signup", json=STAFF).status_code == 409


def test_lockout_after_repeated_failures(app_env):
    app, _ = app_env
    _staff(app)
    c = TestClient(app)
    for _ in range(5):
        assert c.post("/api/auth/login", json={"username": "ed_harbor", "password": "Wrong-Pass-123"}).status_code == 401
    r = c.post("/api/auth/login", json={"username": "ed_harbor", "password": STAFF["password"]})
    assert r.status_code == 429  # locked even with the right password


def test_organizations_cannot_see_each_other(app_env):
    app, _ = app_env
    a = _staff(app)
    _with_plan(app, a)
    a.post("/api/tools/build")
    inst_id = a.get("/api/tools").json()["instruments"][0]["id"]
    sid = a.post("/api/logic-model/sessions", data={"org_name": "Harbor Light Housing"}).json()["session_id"]

    b = _staff(app, username="other_ed", org_name="Other Org")
    assert b.get(f"/api/tools/instruments/{inst_id}").status_code == 404
    assert b.get(f"/api/logic-model/sessions/{sid}").status_code == 404
    assert b.get("/api/tools").json()["instruments"] == []


def test_invite_code_joins_the_same_org(app_env):
    app, _ = app_env
    admin = _staff(app)
    code = admin.post("/api/auth/invite").json()["invite_code"]
    colleague = _staff(app, username="case_mgr", org_name=None, invite_code=code)
    me = colleague.get("/api/auth/me").json()
    assert me["org"]["name"] == "Harbor Light Housing" and me["user"]["role"] == "staff"
    assert colleague.post("/api/auth/invite").status_code == 403  # only admins invite


def test_approving_the_plan_saves_it_to_the_organization(app_env):
    app, _ = app_env
    c = _staff(app)
    sid = c.post("/api/logic-model/sessions", data={"org_name": "Harbor Light Housing"}).json()["session_id"]
    assert c.get("/api/auth/me").json()["studio_session_id"] == sid
    c.post(f"/api/logic-model/sessions/{sid}/intake")
    for n in range(5):
        c.post(f"/api/logic-model/sessions/{sid}/answer", json={"answer": f"answer {n}"})
    c.post(f"/api/logic-model/sessions/{sid}/draft")
    c.post(f"/api/logic-model/sessions/{sid}/approve")
    c.post(f"/api/logic-model/sessions/{sid}/metrics/draft")
    assert c.post(f"/api/logic-model/sessions/{sid}/metrics/approve").json()["phase"] == "complete"
    assert c.get("/api/auth/me").json()["has_plan"] is True


# ============================================================ survey tools, portals, dashboard


def test_full_data_collection_cycle(app_env):
    app, tmp = app_env
    staff = _staff(app)
    _with_plan(app, staff)

    # Tools flow straight from the approved plan's switched-on questions
    tools = staff.post("/api/tools/build").json()
    survey = next(i for i in tools["instruments"] if i["method"] == "participant_survey")
    assessment = next(i for i in tools["instruments"] if i["method"] == "staff_assessment")
    detail = staff.get(f"/api/tools/instruments/{survey['id']}").json()
    texts = [q["text"] for q in detail["questions"]]
    assert "I could handle a major unexpected expense." in texts          # included in the plan
    assert "I am behind on my bills." not in texts                        # switched off in the plan
    assert any(i["name"] == "Families enrolled this month"
               for i in staff.get("/api/tools/indicators").json()["indicators"])  # program_records -> a number

    # Staff refine: hide one question, add an open question, add demographics
    qs = detail["questions"]
    qs[0]["hidden"] = True
    qs.append({"text": "What helped you most?", "scale": "open_text", "waves": ["post"], "source": "staff"})
    detail = staff.put(f"/api/tools/instruments/{survey['id']}",
                       json={"name": detail["name"], "questions": qs, "demographics": ["age_range", "town"]}).json()
    assert detail["questions"][0]["hidden"] and detail["questions"][-1]["source"] == "staff"

    # Survey Review Bot (TestModel offline): the review is stored, nothing is changed automatically
    reviewed = staff.post(f"/api/tools/instruments/{survey['id']}/review").json()
    assert reviewed["review"] is not None and reviewed["questions"] == detail["questions"]

    staff.post(f"/api/tools/instruments/{survey['id']}/publish")
    staff.post(f"/api/tools/instruments/{assessment['id']}/publish")
    people = staff.post("/api/tools/participants", json={"count": 6}).json()["created"]
    assert all(len(p["pin"]) == 6 for p in people)
    assert not any(p["pin"].encode() in (tmp / "test.db").read_bytes() for p in people)  # PINs only hashed

    # Participant portal: each of 6 participants does pre and post
    visible_pre = None
    for n, person in enumerate(people):
        pc = TestClient(app)
        assert pc.post("/api/auth/participant/login", json={"code": person["code"].lower(), "pin": person["pin"]}).status_code == 200
        assert pc.get("/api/tools").status_code == 401  # participants can't reach the staff portal
        mine = pc.get("/api/participant/surveys").json()["surveys"]
        assert [s["id"] for s in mine] == [survey["id"]]  # only participant surveys, not staff assessments
        form = pc.get(f"/api/participant/surveys/{survey['id']}/pre").json()
        visible_pre = form["questions"]
        assert all(q["text"] != qs[0]["text"] for q in visible_pre)  # hidden question not shown
        rated = [q for q in visible_pre if q["scale"] in ("confidence_5", "agree_5")]
        pre = {q["id"]: 2 for q in rated}
        assert pc.post(f"/api/participant/surveys/{survey['id']}/pre",
                       json={"answers": pre, "demographics": {"age_range": "25–34", "town": "New Haven"}}).status_code == 200
        assert pc.post(f"/api/participant/surveys/{survey['id']}/pre", json={"answers": pre}).status_code == 409
        post_form = pc.get(f"/api/participant/surveys/{survey['id']}/post").json()
        post = {q["id"]: (4 if n < 4 else 2) for q in post_form["questions"] if q["scale"] in ("confidence_5", "agree_5")}
        open_q = next(q for q in post_form["questions"] if q["text"] == "What helped you most?")
        post[open_q["id"]] = "=HYPERLINK(\"http://evil\")" if n == 0 else "Having someone check in every week."
        assert pc.post(f"/api/participant/surveys/{survey['id']}/post", json={"answers": post}).status_code == 200

    # A participant with a wrong PIN is refused
    assert TestClient(app).post("/api/auth/participant/login", json={"code": people[0]["code"], "pin": "000000"}).status_code == 401

    # Locked tool: existing questions can't be reworded once answers exist
    d = staff.get(f"/api/tools/instruments/{survey['id']}").json()
    assert d["locked"]
    d["questions"][1]["text"] = "Reworded"
    assert staff.put(f"/api/tools/instruments/{survey['id']}", json={"name": d["name"], "questions": d["questions"],
                                                                     "demographics": d["demographics"]}).status_code == 409

    # Staff bulk entry for the staff assessment, many forms at once
    a_detail = staff.get(f"/api/tools/instruments/{assessment['id']}").json()
    yes_q = [q for q in a_detail["questions"] if q["scale"] == "yes_no"]
    for wave, val in (("pre", 0), ("post", 1)):
        rows = [{"participant_code": p["code"], "answers": {q["id"]: val for q in yes_q}} for p in people]
        rows.append({"participant_code": "NOPE-0000", "answers": {}})
        r = staff.post(f"/api/tools/instruments/{assessment['id']}/responses", json={"wave": wave, "rows": rows}).json()
        assert r["saved"] == 6 and r["errors"][0]["error"].startswith("No participant")

    # Program numbers entered by staff
    ind = staff.get("/api/tools/indicators").json()["indicators"][0]
    staff.post(f"/api/tools/indicators/{ind['id']}/entries",
               json={"entries": [{"period": "2026-09-01", "value": 12}, {"period": "2026-10-01", "value": 15}]})
    staff.post("/api/tools/indicators", json={"name": "Dollars spent on move-in kits", "unit": "$"})

    # Dashboard is populated from the incoming data
    dash = staff.get("/api/dashboard").json()
    assert dash["headline"]["participants"] == 6 and dash["headline"]["matched"] == 6
    improved = next(c for c in dash["charts"] if c["id"] == "improved")
    fc = improved["series"][0]["values"][improved["categories"].index("Financial confidence")]
    assert fc == round(100 * 4 / 6)
    assert improved["series"][0]["values"][improved["categories"].index("Housing stability")] == 100
    age = next(c for c in dash["charts"] if c["id"] == "demo-age_range")
    assert age["categories"] == ["25–34"] and age["series"][0]["values"] == [100] and age["unit"] == "%"
    assert any(c["id"].startswith("indicator-") for c in dash["charts"])
    assert any(a["text"].startswith("Having someone") for a in dash["open_responses"])

    # Downloads: each chart as Excel (with a native chart), all data as Excel, report as PowerPoint
    r = staff.get("/api/dashboard/charts/improved.xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    assert r.status_code == 200 and wb.active._charts
    r = staff.get("/api/dashboard/data.xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    assert "Codebook" in wb.sheetnames and "Program numbers" in wb.sheetnames
    cells = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row]
    assert "'=HYPERLINK(\"http://evil\")" in cells  # formula injection neutralized
    assert not any(isinstance(v, str) and v.startswith("=") for v in cells)
    r = staff.get("/api/dashboard/report.pptx")
    deck = Presentation(io.BytesIO(r.content))
    assert sum(1 for s in deck.slides for sh in s.shapes if sh.has_chart) >= 3

    audit = json.loads((tmp / "audit_log.json").read_text(encoding="utf-8"))
    events = {e["event"] for e in audit}
    assert {"staff_signup", "tools_built", "tool_published", "participants_created", "survey_submitted",
            "responses_entered", "program_numbers_entered", "data_exported"} <= events
    dumped = json.dumps(audit)
    assert all(p["pin"] not in dumped for p in people) and "Harbor-Light-2026" not in dumped


def test_small_demographic_groups_are_hidden(app_env):
    import reports
    from db import Organization, Participant, Response, engine
    from sqlmodel import Session

    app, _ = app_env
    c = _staff(app)
    _with_plan(app, c)
    c.post("/api/tools/build")
    inst_id = next(i["id"] for i in c.get("/api/tools").json()["instruments"] if i["method"] == "participant_survey")
    org_id = c.get("/api/auth/me").json()["org"]["id"]
    with Session(engine()) as db:
        for n in range(7):
            p = Participant(org_id=org_id, code=f"T-{n}", pin_hash="x")
            db.add(p)
            db.commit()
            db.refresh(p)
            db.add(Response(org_id=org_id, instrument_id=inst_id, participant_id=p.id, wave="pre", answers={},
                            demographics={"age_range": "18–24" if n < 5 else "65 or older"}, entered_by="staff"))
        db.commit()
        data = reports.compute(db, db.get(Organization, org_id))
    age = next(ch for ch in data["charts"] if ch["id"] == "demo-age_range")
    assert age["categories"] == ["18–24", reports.HIDDEN]
    assert age["series"][0]["values"] == [71, 29] and age["counts"] == [5, 2]  # % of the 7 who answered


def test_reverse_scoring_flips_rating_and_yes_no_items():
    import surveys

    behind = {"id": "a", "scale": "frequency_5", "reverse": True, "metric": "m", "hidden": False}
    returned = {"id": "b", "scale": "yes_no", "reverse": True, "metric": "n", "hidden": False}
    assert surveys.scored_value(behind, 5) == 1 and surveys.scored_value(behind, 1) == 5
    assert surveys.scored_value(returned, 1) == 0 and surveys.scored_value(returned, 0) == 1
