"""Offline tests for Step 15: program-number templates, percentages, designed reports, Report Designer, privacy."""

import io
import json
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from PIL import Image
from pptx import Presentation
from pydantic_ai.models.test import TestModel

sys.path.insert(0, str(Path(__file__).parent))
from test_prototype import STAFF, _with_plan  # noqa: E402


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setenv("THROUGH_LINE_AUDIT", str(tmp_path / "audit_log.json"))
    monkeypatch.setenv("THROUGH_LINE_SESSIONS", str(tmp_path / "sessions"))
    monkeypatch.setenv("THROUGH_LINE_DB", str(tmp_path / "test.db"))
    monkeypatch.setenv("THROUGH_LINE_BRAND_DIR", str(tmp_path / "brand"))
    from agents import common
    import main
    import security

    security._fails.clear()
    monkeypatch.setattr(common, "MODEL_OVERRIDE", TestModel())
    c = TestClient(main.app)
    assert c.post("/api/auth/signup", json=STAFF).status_code == 200
    _with_plan(main.app, c)
    c.post("/api/tools/build")
    return c, tmp_path


def _logo_png() -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (120, 120), "#1F4E5F").save(out, "PNG")
    return out.getvalue()


def _first_chart_ids(c, n=1):
    return [ch["id"] for ch in c.get("/api/reports").json()["charts"]][:n] or ["improved"]


# ============================================================ program numbers template


def test_numbers_template_round_trip(env):
    c, _ = env
    c.post("/api/tools/indicators", json={"name": "Dollars spent on move-in kits", "unit": "$"})
    r = c.get("/api/tools/indicators/template.xlsx", params={"month": "2026-09"})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Program numbers"]
    rows = [row for row in ws.iter_rows(min_row=5, values_only=True) if row[0]]
    assert rows and all(row[4] == "2026-09" for row in rows)

    for i, row in enumerate(ws.iter_rows(min_row=5)):
        if row[0].value:
            row[5].value = 10 + i
    ws.cell(5, 6).value = "not a number"
    buf = io.BytesIO()
    wb.save(buf)
    r = c.post("/api/tools/indicators/upload", files={"file": ("numbers.xlsx", buf.getvalue(), "application/octet-stream")})
    body = r.json()
    assert r.status_code == 200, r.text
    assert body["saved"] == len(rows) - 1 and len(body["errors"]) == 1 and body["months"] == ["2026-09"]

    # Re-uploading replaces that month instead of duplicating it
    r = c.post("/api/tools/indicators/upload", files={"file": ("numbers.xlsx", buf.getvalue(), "application/octet-stream")})
    entries = [e for i in r.json()["indicators"] for e in i["entries"] if e["period"] == "2026-09-01"]
    assert len(entries) == len(rows) - 1
    # ...and the numbers flow to the dashboard
    assert any(ch["id"].startswith("indicator-") for ch in c.get("/api/dashboard").json()["charts"])


def test_numbers_upload_rejects_other_files(env):
    c, _ = env
    assert c.post("/api/tools/indicators/upload", files={"file": ("x.csv", b"a,b", "text/csv")}).status_code == 400
    empty = io.BytesIO()
    Workbook().save(empty)
    r = c.post("/api/tools/indicators/upload", files={"file": ("x.xlsx", empty.getvalue(), "application/octet-stream")})
    assert r.status_code == 400 and "template" in r.json()["detail"]


# ============================================================ privacy scanner


def test_privacy_scanner_catches_patterns():
    import privacy

    text = "Call Ana at 203-555-0182 or ana@example.org; she lives at 12 Elm Street. Code HLH-4821K, born on May 4, 1990."
    kinds = {f["kind"] for f in privacy.scan(text)}
    assert {"phone number", "email", "street address", "participant code", "date of birth"} <= kinds
    clean, _ = privacy.redact(text)
    assert "203-555" not in clean and "@" not in clean and "HLH-4821K" not in clean
    tree = {"slides": [{"quote": "Reach me at 203 555 0182"}]}
    assert privacy.scan_tree(tree)[0]["where"] == "slides[0].quote"
    assert privacy.scan_tree(privacy.redact_tree(tree)) == []


# ============================================================ quick report + Report Designer


def test_quick_report_is_redesigned(env):
    c, _ = env
    r = c.get("/api/dashboard/report.pptx")
    assert r.status_code == 200 and len(Presentation(io.BytesIO(r.content)).slides) >= 4


def test_report_designer_builds_branded_clean_outputs(env):
    c, tmp = env
    r = c.post("/api/reports/brand",
               data={"standards": "Deep teal and warm gold, friendly and plain.", "primary": "#1F4E5F",
                     "secondary": "#D98E04", "accent": "#F2C14E", "heading_font": "Georgia", "body_font": "Calibri",
                     "tagline": "Home starts here", "contact": "info@example.org"},
               files={"logo": ("logo.png", _logo_png(), "image/png")})
    assert r.status_code == 200, r.text
    assert r.json()["has_logo"] and r.json()["brand"]["primary"] == "#1F4E5F"
    assert c.post("/api/reports/brand", files={"logo": ("x.png", b"not an image", "image/png")}).status_code == 400

    ids = _first_chart_ids(c, 3)
    r = c.post("/api/reports/build", json={"chart_ids": ids, "audience": "Foundation funders", "minutes": 12})
    assert r.status_code == 200, r.text
    run = r.json()
    assert run["privacy"]["clean"] is True
    used = {s.get("chart_id") for s in run["draft"]["slides"]} - {None}
    assert used <= set(ids)  # the agent may only use chosen charts

    deck = c.get(f"/api/reports/{run['id']}/deck.pptx")
    assert deck.status_code == 200 and len(Presentation(io.BytesIO(deck.content)).slides) == len(run["draft"]["slides"])
    pdf = c.get(f"/api/reports/{run['id']}/one-pager.pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")

    log = json.loads((tmp / "audit_log.json").read_text(encoding="utf-8"))
    assert {"brand_saved", "report_built"} <= {e["event"] for e in log}
    finished = [e["agent"] for e in log if e["event"] == "agent_finished"]
    assert finished.count("report_designer") == 1 and finished.count("privacy_reviewer") == 1  # design, then the check


def test_report_designer_scrubs_quotes_before_the_agent(env, monkeypatch):
    """A quote containing contact details never reaches the agent or the final files."""
    import reports
    import routes.reports as rr

    c, _ = env
    original = reports.compute

    def with_bad_quote(db, org):
        d = original(db, org)
        d["open_responses"] = [{"text": "Thank you so much! Text me at 203-555-0182 anytime.", "question": "q",
                                "wave": "post", "instrument": "t", "at": "2026"}]
        return d

    monkeypatch.setattr(reports, "compute", with_bad_quote)
    seen, real = [], rr.run_agent

    async def spy(agent, name, prompt, *a, **k):
        seen.append(prompt)
        return await real(agent, name, prompt, *a, **k)

    monkeypatch.setattr(rr, "run_agent", spy)
    run = c.post("/api/reports/build", json={"chart_ids": _first_chart_ids(c)}).json()
    assert run["privacy"]["quotes_scrubbed_before_agent"] >= 1
    assert seen and all("203-555-0182" not in p for p in seen)
    assert "203-555-0182" not in json.dumps(run["draft"])


def test_reports_are_private_to_the_org(env):
    import main

    c, _ = env
    run = c.post("/api/reports/build", json={"chart_ids": _first_chart_ids(c)}).json()
    other = TestClient(main.app)
    other.post("/api/auth/signup", json={**STAFF, "username": "someone_else", "org_name": "Other Org"})
    assert other.get(f"/api/reports/{run['id']}").status_code == 404
    assert other.get(f"/api/reports/{run['id']}/deck.pptx").status_code == 404
    assert TestClient(main.app).get("/api/reports").status_code == 401
