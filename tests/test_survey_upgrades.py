"""Multiple-choice questions, the timing check, review requests, and entering responses from Excel."""

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def _tool(c):
    return c.get("/api/tools").json()["instruments"][0]["id"]


def _with_choice(c, iid):
    d = c.get(f"/api/tools/instruments/{iid}").json()
    qs = d["questions"] + [{"id": "q_barriers", "text": "Which of these got in the way? Choose all that apply.",
                            "scale": "multi_choice", "options": ["Transportation", "Childcare", "Cost, fees", "Other", ""],
                            "waves": ["pre", "post"], "source": "staff"},
                           {"id": "q_heard", "text": "How did you hear about us?", "scale": "choice",
                            "options": ["Friend", "School", "Other"], "waves": ["pre"], "source": "staff"}]
    r = c.put(f"/api/tools/instruments/{iid}", json={"name": d["name"], "questions": qs, "demographics": ["town"]})
    assert r.status_code == 200, r.text
    return r.json()


def test_choice_questions_validate_and_count(org_env):
    c, _ = org_env
    iid = _tool(c)
    d = _with_choice(c, iid)
    q = next(q for q in d["questions"] if q["id"] == "q_barriers")
    assert q["options"] == ["Transportation", "Childcare", "Cost, fees", "Other"]  # blanks dropped, commas kept

    bad = c.put(f"/api/tools/instruments/{iid}", json={"name": d["name"], "demographics": [], "questions": [
        *d["questions"], {"text": "Pick one", "scale": "choice", "options": ["Only one"], "waves": ["post"]}]})
    assert bad.status_code == 400 and "2 and 12" in bad.text

    codes = [p["code"] for p in c.post("/api/tools/participants", json={"count": 6}).json()["created"]]
    rows = [{"participant_code": code, "answers": {"q_barriers": "Transportation; Cost, fees" if i % 2 else "Childcare",
                                                 "q_heard": "School"}} for i, code in enumerate(codes)]
    r = c.post(f"/api/tools/instruments/{iid}/responses", json={"wave": "pre", "rows": rows})
    assert r.json()["saved"] == 6, r.text
    r = c.post(f"/api/tools/instruments/{iid}/responses",
               json={"wave": "pre", "rows": [{"participant_code": codes[0], "answers": {"q_heard": "Billboard"}}]})
    assert r.json()["errors"] and "choices" in r.json()["errors"][0]["error"]

    charts = {ch["id"]: ch for ch in c.get("/api/dashboard").json()["charts"]}
    barriers = charts[f"choice-{iid}-q_barriers"]
    assert barriers["group"] == "answers" and barriers["series"][0]["values"] == [50, 50, 50, 0]  # 3 of 6 picked each of the first three
    assert c.get("/api/dashboard/data.xlsx").status_code == 200  # lists export cleanly

    # Once answers exist, choices can be added but not removed or renamed.
    d = c.get(f"/api/tools/instruments/{iid}").json()
    for q in d["questions"]:
        if q["id"] == "q_heard":
            q["options"] = ["Friend", "School"]
    assert c.put(f"/api/tools/instruments/{iid}", json={"name": d["name"], "questions": d["questions"],
                                                         "demographics": d["demographics"]}).status_code == 409


def test_pre_questions_that_assume_the_program_are_flagged(org_env):
    import surveys

    q = lambda text, waves, scale="agree_5": {"id": text[:5], "text": text, "scale": scale, "waves": waves,
                                              "hidden": False, "metric": None}
    issues = surveys.precheck([q("Since you started the program, I feel more confident", ["pre", "post"]),
                               q("What has changed for you?", ["pre", "post"], "open_text"),
                               q("I feel confident managing my budget", ["pre", "post"]),
                               q("How much has the program helped you?", ["post"])], "staff_assessment")
    timing = [i for i in issues if "won't make sense at pre" in i["problem"]]
    assert {i["question_id"] for i in timing} == {"Since", "What "}
    assert next(i for i in timing if i["question_id"] == "What ")["fix"] == {"action": "set_waves", "waves": ["post"]}


def test_review_bot_receives_the_staff_request(org_env, monkeypatch):
    import routes.tools as rt

    c, _ = org_env
    seen = []
    real = rt.run_agent

    async def spy(agent, name, prompt, *a, **k):
        seen.append(prompt)
        return await real(agent, name, prompt, *a, **k)

    monkeypatch.setattr(rt, "run_agent", spy)
    r = c.post(f"/api/tools/instruments/{_tool(c)}/review", json={"request": "Turn the written questions into choices"})
    assert r.status_code == 200 and '"staff_request": "Turn the written questions into choices"' in seen[0]
    assert c.post(f"/api/tools/instruments/{_tool(c)}/review").status_code == 200  # the request is optional


def test_responses_spreadsheet_round_trip(org_env):
    c, _ = org_env
    iid = _tool(c)
    d = _with_choice(c, iid)
    codes = [p["code"] for p in c.post("/api/tools/participants", json={"count": 3}).json()["created"]]
    r = c.get(f"/api/tools/instruments/{iid}/responses/template.xlsx", params={"wave": "pre"})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Responses"]
    ids = [cell.value for cell in ws[3]][1:]
    assert ws.row_dimensions[3].hidden and [row[0].value for row in ws.iter_rows(min_row=5)] == sorted(codes)

    rated = next(q for q in d["questions"] if q["scale"] == "agree_5" and "pre" in q["waves"] and not q["hidden"])
    col = {qid: i + 2 for i, qid in enumerate(ids)}
    ws.cell(5, col[rated["id"]]).value = "Agree"  # a label from the dropdown
    ws.cell(5, col["q_barriers"]).value = "Childcare; Other"
    ws.cell(5, col["demo:town"]).value = "New Haven"
    ws.cell(6, col["q_heard"]).value = "Not a choice"
    buf = io.BytesIO()
    wb.save(buf)
    r = c.post(f"/api/tools/instruments/{iid}/responses/upload", params={"wave": "pre"},
               files={"file": ("pre.xlsx", buf.getvalue(), "application/octet-stream")})
    body = r.json()
    assert r.status_code == 200 and body["saved"] == 1 and len(body["errors"]) == 1, body

    # The answers were stored as values (Agree = 4) and come back pre-filled as labels.
    wb2 = load_workbook(io.BytesIO(c.get(f"/api/tools/instruments/{iid}/responses/template.xlsx",
                                         params={"wave": "pre"}).content))
    row = next(r for r in wb2["Responses"].iter_rows(min_row=5, values_only=True) if r[0] == sorted(codes)[0])
    assert row[col[rated["id"]] - 1] == "Agree" and row[col["q_barriers"] - 1] == "Childcare; Other"

    assert c.post(f"/api/tools/instruments/{iid}/responses/upload", params={"wave": "pre"},
                  files={"file": ("x.csv", b"a,b", "text/csv")}).status_code == 400


def test_responses_are_private_to_the_org(org_env):
    import main

    c, _ = org_env
    iid = _tool(c)
    other = TestClient(main.app)
    other.post("/api/auth/signup", json={"username": "other_org_ed", "password": "Harbor-Light-2026",
                                         "display_name": "X", "org_name": "Other Org"})
    assert other.get(f"/api/tools/instruments/{iid}/responses/template.xlsx", params={"wave": "pre"}).status_code == 404
